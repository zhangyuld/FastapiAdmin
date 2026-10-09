import re
from io import BytesIO
from urllib.parse import unquote

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _customer_payload(customer_code: str) -> dict:
    """构造客户创建请求。"""
    return {
        "customer_code": customer_code,
        "customer_name": "测试科技有限公司",
        "contact": "测试联系人",
        "phone": "13800138000",
        "address": "测试联系地址",
        "remark": "客户接口测试",
        "status": 0,
    }


def test_customer_crud_search_status_and_export(
    test_client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    """客户支持 CRUD、模糊筛选、审计筛选、状态修改和字段导出。"""
    payload = _customer_payload("TEST-CUSTOMER-CRUD-001")
    created = test_client.post(
        "/inventory/basic/customer/create",
        headers=auth_headers,
        json=payload,
    )
    assert created.status_code == 200, created.text
    created_data = created.json()["data"]
    customer_id = created_data["id"]
    assert created_data["customer_code"] == payload["customer_code"]
    assert created_data["created_id"] is not None

    listed = test_client.get(
        "/inventory/basic/customer/list",
        headers=auth_headers,
        params={
            "page_no": 1,
            "page_size": 10,
            "customer_code": "CUSTOMER-CRUD",
            "customer_name": "科技",
        },
    )
    assert listed.status_code == 200, listed.text
    assert any(item["id"] == customer_id for item in listed.json()["data"]["items"])

    audit_listed = test_client.get(
        "/inventory/basic/customer/list",
        headers=auth_headers,
        params=[
            ("page_no", "1"),
            ("page_size", "10"),
            ("created_id", str(created_data["created_id"])),
            ("created_time", "2020-01-01 00:00:00"),
            ("created_time", "2030-01-01 00:00:00"),
        ],
    )
    assert audit_listed.status_code == 200, audit_listed.text
    assert any(item["id"] == customer_id for item in audit_listed.json()["data"]["items"])

    update_payload = {
        **payload,
        "customer_name": "测试科技客户（已修改）",
        "phone": "12345678901234567890123456789012",
        "address": "测" * 300,
    }
    updated = test_client.post(
        f"/inventory/basic/customer/update/{customer_id}",
        headers=auth_headers,
        json=update_payload,
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["data"]["customer_name"] == update_payload["customer_name"]
    assert updated.json()["data"]["address"] == update_payload["address"]

    disabled = test_client.post(
        "/inventory/basic/customer/status/batch",
        headers=auth_headers,
        json={"ids": [customer_id, customer_id], "status": 1},
    )
    assert disabled.status_code == 200, disabled.text
    detail = test_client.get(
        f"/inventory/basic/customer/detail/{customer_id}",
        headers=auth_headers,
    )
    assert detail.status_code == 200, detail.text
    assert detail.json()["data"]["status"] == 1

    exported = test_client.post(
        "/inventory/basic/customer/export",
        headers=auth_headers,
        json={
            "query": {"customer_code": "CUSTOMER-CRUD"},
            "fields": ["customer_code", "customer_name", "status"],
            "format": "xlsx",
            "sheet_name": "客户测试",
            "filename": "客户自定义清单.xlsx",
        },
    )
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    content_disposition = unquote(exported.headers["content-disposition"])
    assert "filename*=UTF-8''客户自定义清单.xlsx" in content_disposition
    workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
    worksheet = workbook.active
    assert worksheet is not None
    assert worksheet.title == "客户测试"
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    assert list(rows[0]) == ["客户编码", "客户名称", "状态"]
    assert any(row[0] == payload["customer_code"] and row[2] == "停用" for row in rows[1:])

    exported_csv = test_client.post(
        "/inventory/basic/customer/export",
        headers=auth_headers,
        json={
            "query": {"customer_code": "CUSTOMER-CRUD"},
            "fields": ["customer_code", "phone"],
            "format": "csv",
            "sheet_name": "客户测试",
            "filename": "   ",
        },
    )
    assert exported_csv.status_code == 200, exported_csv.text
    csv_content_disposition = unquote(exported_csv.headers["content-disposition"])
    assert re.search(
        r"filename\*=UTF-8''客户列表_\d{8}_\d{6}\.csv",
        csv_content_disposition,
    )
    assert "客户编码" in exported_csv.content.decode("utf-8-sig")

    deleted = test_client.post(
        "/inventory/basic/customer/delete",
        headers=auth_headers,
        json={"ids": [customer_id, customer_id]},
    )
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["data"] == [customer_id]


def test_customer_business_validation(
    test_client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    """客户拒绝重复编码、非法字段、非法批量参数和非法工作表名称。"""
    payload = _customer_payload("TEST-CUSTOMER-VALIDATE-001")
    created = test_client.post(
        "/inventory/basic/customer/create",
        headers=auth_headers,
        json=payload,
    )
    assert created.status_code == 200, created.text
    customer_id = created.json()["data"]["id"]

    duplicate = test_client.post(
        "/inventory/basic/customer/create",
        headers=auth_headers,
        json=payload,
    )
    assert duplicate.status_code == 409, duplicate.text
    assert "客户编码已存在" in duplicate.json()["msg"]

    invalid_contact = test_client.post(
        "/inventory/basic/customer/create",
        headers=auth_headers,
        json={**_customer_payload("TEST-CUSTOMER-CONTACT"), "contact": "联" * 65},
    )
    assert invalid_contact.status_code == 422, invalid_contact.text

    empty_status = test_client.post(
        "/inventory/basic/customer/status/batch",
        headers=auth_headers,
        json={"ids": [], "status": 1},
    )
    assert empty_status.status_code == 422, empty_status.text

    missing_status = test_client.post(
        "/inventory/basic/customer/status/batch",
        headers=auth_headers,
        json={"ids": [999999], "status": 1},
    )
    assert missing_status.status_code == 404, missing_status.text
    assert "客户不存在" in missing_status.json()["msg"]

    invalid_sheet_name = test_client.post(
        "/inventory/basic/customer/export",
        headers=auth_headers,
        json={
            "query": {},
            "fields": ["customer_code"],
            "format": "xlsx",
            "sheet_name": "客户/档案",
        },
    )
    assert invalid_sheet_name.status_code == 422, invalid_sheet_name.text

    invalid_filename = test_client.post(
        "/inventory/basic/customer/export",
        headers=auth_headers,
        json={
            "query": {},
            "fields": ["customer_code"],
            "format": "xlsx",
            "filename": "客户/档案",
        },
    )
    assert invalid_filename.status_code == 422, invalid_filename.text

    deleted = test_client.post(
        "/inventory/basic/customer/delete",
        headers=auth_headers,
        json={"ids": [customer_id]},
    )
    assert deleted.status_code == 200, deleted.text

    recreate_deleted_code = test_client.post(
        "/inventory/basic/customer/create",
        headers=auth_headers,
        json=payload,
    )
    assert recreate_deleted_code.status_code == 409, recreate_deleted_code.text
    assert "客户编码已存在" in recreate_deleted_code.json()["msg"]
