from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _create_category(
    client: TestClient,
    headers: dict[str, str],
    *,
    name: str,
    status: int = 0,
) -> dict:
    """通过接口创建商品测试分类。"""
    response = client.post(
        "/inventory/basic/category/create",
        headers=headers,
        json={"name": name, "parent_id": None, "sort": 1, "status": status},
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _goods_payload(category_id: int, goods_code: str) -> dict:
    """构造商品创建请求。"""
    return {
        "goods_name": "测试商品",
        "goods_code": goods_code,
        "category_id": category_id,
        "spec": None,
        "unit": "件",
        "reference_cost_price": "3.1000",
        "sale_price": "5.2000",
        "stock_min": "2.000",
        "stock_max": "0",
        "remark": None,
        "status": 0,
    }


def test_goods_create_and_list_with_category_name(
    test_client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    """商品支持创建、精确分类筛选并返回分类名称。"""
    category = _create_category(test_client, auth_headers, name="商品接口测试分类")
    payload = _goods_payload(category["id"], "TEST-GOODS-001")

    created = test_client.post(
        "/inventory/basic/goods/create",
        headers=auth_headers,
        json=payload,
    )
    assert created.status_code == 200, created.text
    created_data = created.json()["data"]
    assert created_data["goods_code"] == payload["goods_code"]
    assert created_data["category_name"] == category["name"]
    assert created_data["spec"] is None

    listed = test_client.get(
        "/inventory/basic/goods/list",
        headers=auth_headers,
        params={"page_no": 1, "page_size": 10, "category_id": category["id"], "goods_code": "GOODS-001"},
    )
    assert listed.status_code == 200, listed.text
    item = next(item for item in listed.json()["data"]["items"] if item["id"] == created_data["id"])
    assert item["category_name"] == category["name"]
    assert item["spec"] is None

    exported = test_client.post(
        "/inventory/basic/goods/export",
        headers=auth_headers,
        json={
            "query": {},
            "fields": ["goods_code", "category_name"],
            "format": "xlsx",
            "sheet_name": "商品测试",
        },
    )
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "goods.xlsx" in exported.headers["content-disposition"]
    workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
    worksheet = workbook.active
    assert worksheet is not None
    assert worksheet.title == "商品测试"
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    headers = list(rows[0])
    assert headers == ["商品编码", "商品分类"]
    goods_code_index = headers.index("商品编码")
    category_name_index = headers.index("商品分类")
    exported_row = next(row for row in rows[1:] if row[goods_code_index] == payload["goods_code"])
    assert exported_row[category_name_index] == category["name"]

    deleted = test_client.post(
        "/inventory/basic/goods/delete",
        headers=auth_headers,
        json={"ids": [created_data["id"], created_data["id"]]},
    )
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["data"] == [created_data["id"]]

    listed_after_delete = test_client.get(
        "/inventory/basic/goods/list",
        headers=auth_headers,
        params={"page_no": 1, "page_size": 10, "category_id": category["id"]},
    )
    assert listed_after_delete.status_code == 200, listed_after_delete.text
    assert all(item["id"] != created_data["id"] for item in listed_after_delete.json()["data"]["items"])


def test_goods_create_business_validation(
    test_client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    """商品创建拒绝重复编码、无效分类和不合法价格库存。"""
    category = _create_category(test_client, auth_headers, name="商品校验测试分类")
    payload = _goods_payload(category["id"], "TEST-GOODS-VALIDATE")
    first = test_client.post("/inventory/basic/goods/create", headers=auth_headers, json=payload)
    assert first.status_code == 200, first.text

    duplicate = test_client.post("/inventory/basic/goods/create", headers=auth_headers, json=payload)
    assert duplicate.status_code == 409
    assert "商品编码已存在" in duplicate.json()["msg"]

    missing_category_payload = _goods_payload(999999, "TEST-GOODS-NO-CATEGORY")
    missing_category = test_client.post(
        "/inventory/basic/goods/create",
        headers=auth_headers,
        json=missing_category_payload,
    )
    assert missing_category.status_code == 404
    assert "商品分类不存在" in missing_category.json()["msg"]

    empty_delete = test_client.post(
        "/inventory/basic/goods/delete",
        headers=auth_headers,
        json={"ids": []},
    )
    assert empty_delete.status_code == 422

    missing_goods = test_client.post(
        "/inventory/basic/goods/delete",
        headers=auth_headers,
        json={"ids": [999999]},
    )
    assert missing_goods.status_code == 404
    assert "商品不存在" in missing_goods.json()["msg"]

    disabled_category = _create_category(test_client, auth_headers, name="商品停用分类", status=1)
    disabled_category_payload = _goods_payload(disabled_category["id"], "TEST-GOODS-DISABLED-CATEGORY")
    disabled_category_result = test_client.post(
        "/inventory/basic/goods/create",
        headers=auth_headers,
        json=disabled_category_payload,
    )
    assert disabled_category_result.status_code == 400
    assert "商品分类已停用" in disabled_category_result.json()["msg"]

    invalid_price_payload = _goods_payload(category["id"], "TEST-GOODS-PRICE")
    invalid_price_payload["sale_price"] = "1.0000"
    invalid_price = test_client.post(
        "/inventory/basic/goods/create",
        headers=auth_headers,
        json=invalid_price_payload,
    )
    assert invalid_price.status_code == 422

    invalid_stock_payload = _goods_payload(category["id"], "TEST-GOODS-STOCK")
    invalid_stock_payload["stock_max"] = "1.000"
    invalid_stock = test_client.post(
        "/inventory/basic/goods/create",
        headers=auth_headers,
        json=invalid_stock_payload,
    )
    assert invalid_stock.status_code == 422
