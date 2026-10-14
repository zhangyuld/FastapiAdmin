"""库存报表余额、预警、期间反向流水、独立权限及字段导出回归。"""

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.database import async_db_session
from app.modules.inventory.basic.model import GoodsCategoryModel, GoodsModel, WarehouseModel
from app.modules.inventory.stock.model import StockBalanceModel, StockLogModel
from app.modules.system.user.model import UserModel

PREFIX = "/inventory/report/stock"


@pytest.fixture
def stock_report_data(test_client, auth_headers):
    """临时库建立余额和不可变流水，不向开发库写入测试单据。"""
    key = uuid4().hex[:12]

    async def seed():
        async with async_db_session() as db, db.begin():
            user_id = await db.scalar(select(UserModel.id).where(UserModel.username == "super"))
            categories = [GoodsCategoryModel(name=key + "Category" + str(i)) for i in range(2)]
            warehouses = [WarehouseModel(warehouse_code=key + str(i), warehouse_name=key + "Warehouse" + str(i)) for i in range(2)]
            db.add_all([*categories, *warehouses])
            await db.flush()
            goods = []
            limits = [("5", "20"), ("2", "10"), ("0", "0"), ("5", "20"), ("5", "20"), ("0", "0"), ("0", "0"), ("0", "0")]
            for i, (lower, upper) in enumerate(limits):
                goods.append(
                    GoodsModel(
                        goods_code=key + ("%_G1" if i == 1 else "G" + str(i)),
                        goods_name=key + "Goods" + str(i),
                        category_id=categories[i % 2].id,
                        spec="current",
                        unit="箱" if i == 1 else "件",
                        stock_min=Decimal(lower),
                        stock_max=Decimal(upper),
                    )
                )
            db.add_all(goods)
            await db.flush()
            for warehouse, g, quantity, amount in [
                (0, 0, "10", "50"),
                (1, 0, "0", "0"),
                (0, 1, "12", "24"),
                (1, 1, "2", "4"),
                (0, 2, "100", "0"),
                (0, 3, "5", "5"),
                (0, 4, "20", "20"),
                (0, 6, "0", "0"),
            ]:
                q, a = Decimal(quantity), Decimal(amount)
                db.add(
                    StockBalanceModel(
                        warehouse_id=warehouses[warehouse].id,
                        goods_id=goods[g].id,
                        quantity=q,
                        stock_amount=a,
                        avg_cost=a / q if q else Decimal("0"),
                        created_id=user_id,
                    )
                )
            db.add(StockBalanceModel(warehouse_id=warehouses[0].id, goods_id=goods[7].id, quantity=999, avg_cost=1, stock_amount=999, is_deleted=True))
            index = max(1000000, await db.scalar(select(func.max(StockLogModel.business_item_id))) or 0) + 100

            async def log(warehouse, g, kind, quantity, amount, day="2026-10-01", operation="confirm", reversal=None, deleted=False):
                """流水方向由带符号数量和成本决定，作废仍使用原业务日期。"""
                nonlocal index
                index += 1
                q = Decimal(quantity)
                before = Decimal("100")
                row = StockLogModel(
                    warehouse_id=warehouses[warehouse].id,
                    goods_id=goods[g].id,
                    business_type=kind,
                    operation_type=operation,
                    business_id=index,
                    business_item_id=index,
                    business_no=key + str(index),
                    business_date=date.fromisoformat(day),
                    change_quantity=q,
                    before_quantity=before,
                    after_quantity=before + q,
                    unit_cost=Decimal("5"),
                    cost_amount=Decimal(amount),
                    reversal_log_id=reversal.id if reversal else None,
                    posted_time=datetime(2026, 10, 20 if operation == "void" else 1, tzinfo=UTC),
                    created_id=user_id,
                    is_deleted=deleted,
                )
                db.add(row)
                await db.flush()
                return row

            await log(0, 0, "opening_balance", "10", "50")
            await log(0, 0, "purchase_in", "8", "40", "2026-10-02")
            await log(0, 0, "sale_out", "-4", "-20", "2026-10-03")
            await log(0, 0, "sale_return", "1", "5", "2026-10-04")
            await log(0, 0, "purchase_return", "-2", "-10", "2026-10-05")
            await log(0, 0, "stock_check_gain", "2", "10", "2026-10-06")
            await log(0, 0, "stock_check_loss", "-1", "-5", "2026-10-07")
            incoming = await log(0, 0, "purchase_in", "3", "15", "2026-09-30")
            await log(0, 0, "purchase_in", "-3", "-15", "2026-09-30", "void", incoming)
            outgoing = await log(0, 0, "sale_out", "-2", "-10", "2026-09-30")
            await log(0, 0, "sale_out", "2", "10", "2026-09-30", "void", outgoing)
            await log(0, 1, "purchase_in", "5", "0")
            await log(0, 1, "sale_out", "-2", "0")
            await log(1, 0, "sale_out", "-1", "-5")
            await log(0, 2, "stock_check_gain", "0", "0.01")
            await log(0, 0, "purchase_in", "1", "5", "2026-11-01")
            await log(0, 0, "purchase_in", "99", "999", deleted=True)
            goods[0].status = 1
            goods[0].is_deleted = True
            warehouses[0].status = 1
            warehouses[0].is_deleted = True
            return {
                "key": key,
                "warehouses": [row.id for row in warehouses],
                "categories": [row.id for row in categories],
                "goods": [row.id for row in goods],
                "codes": [row.goods_code for row in goods],
            }

    with patch("app.core.router_class._write_operation_log_async", new=AsyncMock()), patch("app.modules.system.auth.service._write_login_log", new=AsyncMock()):
        yield test_client.portal.call(seed)


def data(client, headers, fixture, endpoint, params=None):
    """限定测试仓库，分页、汇总及不同日期查询均不串入其他测试数据。"""
    response = client.get(PREFIX + endpoint, headers=headers, params={"warehouse_name": fixture["key"], **(params or {})})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_balance_uses_existing_records_and_distinct_goods(test_client, auth_headers, stock_report_data):
    """零余额保留，未建账组合不生成，跨仓商品数去重，删除余额不计入。"""
    summary = data(test_client, auth_headers, stock_report_data, "/summary")
    assert summary == {"goods_count": 6, "stock_amount": "103.00"}
    result = data(test_client, auth_headers, stock_report_data, "/list")
    assert result["total"] == 8 and len({row["id"] for row in result["items"]}) == 8
    rows = result["items"]
    assert all("source" not in row and "in_amount" not in row for row in rows)
    assert len([row for row in rows if row["quantity"] == "0.000"]) == 2
    assert all(row["goods_code"] not in stock_report_data["codes"][5::2] for row in rows)
    row = next(row for row in rows if row["goods_code"] == stock_report_data["codes"][0] and row["quantity"] == "10.000")
    assert row["avg_cost"] == "5.0000" and row["stock_amount"] == "50.00"


def test_warning_boundaries_and_unlimited_upper(test_client, auth_headers, stock_report_data):
    """等于上下限是正常，上限零不限，零余额依然按下限判断。"""
    for warning, count, amount in [("low", 1, "0.00"), ("high", 1, "24.00"), ("normal", 6, "79.00")]:
        result = data(test_client, auth_headers, stock_report_data, "/list", {"warning": warning})
        assert result["total"] == count
        assert all(row["warning"] == warning for row in result["items"])
        assert data(test_client, auth_headers, stock_report_data, "/summary", {"warning": warning})["stock_amount"] == amount
    result = data(test_client, auth_headers, stock_report_data, "/list", {"goods_code": stock_report_data["codes"][2]})
    assert result["items"][0]["quantity"] == "100.000" and result["items"][0]["warning"] == "normal"


def test_movement_includes_all_business_directions_and_cost(test_client, auth_headers, stock_report_data):
    """期初、盘盈盘亏、退货及反向流水都按带符号成本统计，不取参考进价。"""
    query = {"view": "movement", "business_date": ["2026-09-01", "2026-10-31"]}
    summary = data(test_client, auth_headers, stock_report_data, "/summary", query)
    assert summary == {"goods_count": 3, "in_amount": "130.01", "out_amount": "65.00", "net_amount": "65.01"}
    result = data(test_client, auth_headers, stock_report_data, "/list", query)
    assert result["total"] == 4
    main = next(row for row in result["items"] if row["goods_code"] == stock_report_data["codes"][0] and row["warehouse_name"].endswith("0"))
    assert [main[field] for field in ("in_quantity", "out_quantity", "net_quantity")] == ["26.000", "12.000", "14.000"]
    assert [main[field] for field in ("in_amount", "out_amount", "net_amount")] == ["130.00", "60.00", "70.00"]
    assert all("source" not in row and "stock_amount" not in row for row in result["items"])
    assert "net_quantity" not in summary


def test_void_retains_gross_flows_on_original_business_date(test_client, auth_headers, stock_report_data):
    """作废入账在十月但业务日期仍九月，原期间入出各计一次且净额抵消。"""
    query = {"view": "movement", "business_date": ["2026-09-30", "2026-09-30"]}
    summary = data(test_client, auth_headers, stock_report_data, "/summary", query)
    assert summary == {"goods_count": 1, "in_amount": "25.00", "out_amount": "25.00", "net_amount": "0.00"}
    row = data(test_client, auth_headers, stock_report_data, "/list", query)["items"][0]
    assert row["in_quantity"] == row["out_quantity"] == "5.000" and row["net_quantity"] == "0.000"
    october = data(test_client, auth_headers, stock_report_data, "/summary", {"view": "movement", "business_date": ["2026-10-01", "2026-10-31"]})
    assert october["in_amount"] == "105.01" and october["out_amount"] == "40.00" and october["net_amount"] == "65.01"
    all_time = data(test_client, auth_headers, stock_report_data, "/summary", {"view": "movement"})
    assert all_time["in_amount"] == "135.01" and all_time["net_amount"] == "70.01"


def test_free_goods_zero_quantity_cost_and_negative_net(test_client, auth_headers, stock_report_data):
    """免费商品不丢数量，零数量成本不丢金额，出库期间允许负净变动。"""
    result = data(test_client, auth_headers, stock_report_data, "/list", {"view": "movement", "goods_code": stock_report_data["codes"][1]})
    row = result["items"][0]
    assert (row["in_quantity"], row["out_quantity"], row["net_quantity"], row["net_amount"]) == ("5.000", "2.000", "3.000", "0.00")
    cost_only = data(test_client, auth_headers, stock_report_data, "/list", {"view": "movement", "goods_code": stock_report_data["codes"][2]})["items"][0]
    assert cost_only["net_quantity"] == "0.000" and cost_only["net_amount"] == "0.01"
    negative = data(test_client, auth_headers, stock_report_data, "/list", {"view": "movement", "warehouse_id": stock_report_data["warehouses"][1]})["items"][0]
    assert negative["net_quantity"] == "-1.000" and negative["net_amount"] == "-5.00"


def test_current_archives_filters_and_literal_wildcards(test_client, auth_headers, stock_report_data):
    """流水没有商品快照，改名后使用当前档案，停用或删除档案仍可查询。"""

    async def rename():
        async with async_db_session() as db, db.begin():
            row = await db.get(GoodsModel, stock_report_data["goods"][0])
            row.goods_code = stock_report_data["key"] + "CURRENT"
            row.goods_name = stock_report_data["key"] + "Renamed"

    test_client.portal.call(rename)
    assert data(test_client, auth_headers, stock_report_data, "/summary", {"goods_code": stock_report_data["codes"][0]})["goods_count"] == 0
    assert data(test_client, auth_headers, stock_report_data, "/summary", {"view": "movement", "goods_name": stock_report_data["key"] + "Renamed"})["net_amount"] == "70.00"
    assert data(test_client, auth_headers, stock_report_data, "/summary", {"goods_code": "%_"}) == {"goods_count": 1, "stock_amount": "28.00"}
    expected = {"goods_count": 2, "stock_amount": "33.00"}
    assert data(test_client, auth_headers, stock_report_data, "/summary", {"category_id": stock_report_data["categories"][1]}) == expected
    assert data(test_client, auth_headers, stock_report_data, "/summary", {"category_name": stock_report_data["key"] + "Category1"}) == expected


@pytest.mark.parametrize("view", ["balance", "movement"])
def test_stable_paging_and_full_summary(test_client, auth_headers, stock_report_data, view):
    """列表分组后分页，稳定键跨页不重复，汇总不受分页字段影响。"""
    result = data(test_client, auth_headers, stock_report_data, "/list", {"view": view})
    ids = [data(test_client, auth_headers, stock_report_data, "/list", {"view": view, "page_no": i, "page_size": 1})["items"][0]["id"] for i in range(1, result["total"] + 1)]
    assert len(ids) == len(set(ids)) and set(ids) == {row["id"] for row in result["items"]}
    assert data(test_client, auth_headers, stock_report_data, "/summary", {"view": view, "page_no": 9, "page_size": 1}) == data(
        test_client, auth_headers, stock_report_data, "/summary", {"view": view}
    )
    ordered = data(test_client, auth_headers, stock_report_data, "/list", {"view": view, "order_by": '[{"goods_code":"desc"}]'})
    assert [row["goods_code"] for row in ordered["items"]] == sorted([row["goods_code"] for row in ordered["items"]], reverse=True)


@pytest.mark.parametrize(
    "params",
    [
        {"view": "goods"},
        {"warehouse_id": 0},
        {"category_id": -1},
        {"business_date": ["2026-10-01", "2026-10-31"]},
        {"view": "movement", "warning": "low"},
        {"warning": "unknown"},
        {"view": "movement", "business_date": ["2026-10-02", "2026-10-01"]},
        {"view": "movement", "business_date": ["2026-10-01"]},
        {"order_by": '[{"stock_amount":"drop"}]'},
        {"view": "balance", "order_by": '[{"in_amount":"asc"}]'},
        {"view": "movement", "order_by": '[{"warning":"asc"}]'},
        {"order_by": "[{}]"},
    ],
)
def test_invalid_query_rejected(test_client, auth_headers, stock_report_data, params):
    """后端拒绝不合法视图、区间、跨视图预警和非白名单排序。"""
    response = test_client.get(PREFIX + "/list", headers=auth_headers, params={"warehouse_name": stock_report_data["key"], **params})
    assert response.status_code == 422, response.text


@pytest.mark.parametrize(
    "view,expected",
    [
        ("balance", {"goods_count": 0, "stock_amount": "0.00"}),
        ("movement", {"goods_count": 0, "in_amount": "0.00", "out_amount": "0.00", "net_amount": "0.00"}),
    ],
)
def test_empty_results_have_zero_summary(test_client, auth_headers, stock_report_data, view, expected):
    """不存在的筛选返回空列表与真实零汇总，不返回缺失或模拟金额。"""
    query = {"view": view, "goods_code": "unmatched"}
    assert data(test_client, auth_headers, stock_report_data, "/list", query)["items"] == []
    assert data(test_client, auth_headers, stock_report_data, "/summary", query) == expected


def test_exports_follow_view_fields_order_and_precision(test_client, auth_headers, stock_report_data):
    """CSV/Excel 共用筛选和字段顺序，预警中文化，数量和平均成本保持精度。"""
    query = {"warehouse_name": stock_report_data["key"], "warning": "low"}
    body = {"query": query, "fields": ["warning", "avg_cost", "quantity", "goods_code"], "format": "csv", "filename": "库存核对.xlsx"}
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "预警,平均成本,库存数量,商品编码"
    assert lines[1] == "低于下限,0.0000,0.000," + stock_report_data["codes"][0]
    assert "%E5%BA%93%E5%AD%98%E6%A0%B8%E5%AF%B9.csv" in response.headers["content-disposition"]
    body.update(query={"view": "movement", "warehouse_id": stock_report_data["warehouses"][1]}, fields=["net_amount", "goods_name", "net_quantity"], format="xlsx", filename="", sheet_name="库存核对")
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "库存核对" and sheet.max_row == 2 and sheet.max_column == 3
    assert [cell.value for cell in sheet[1]] == ["净变动金额", "商品名称", "净变动数量"]
    assert sheet.cell(2, 1).value == "-5.00" and sheet.cell(2, 3).value == "-1.000"
    assert "%E5%BA%93%E5%AD%98%E6%8A%A5%E8%A1%A8_" in response.headers["content-disposition"]
    body.update(query={"view": "movement", "warehouse_name": stock_report_data["key"]}, fields=["warning"])
    assert test_client.post(PREFIX + "/export", headers=auth_headers, json=body).status_code == 422
    body.update(query={"warehouse_name": stock_report_data["key"]}, fields=["in_amount"])
    assert test_client.post(PREFIX + "/export", headers=auth_headers, json=body).status_code == 422


@pytest.mark.parametrize(
    "options",
    [
        {"fields": []},
        {"fields": ["not_a_column"]},
        {"format": "pdf"},
        {"filename": "../bad"},
        {"filename": "a\nb"},
        {"sheet_name": "bad[]"},
        {"sheet_name": "x" * 32},
        {"query": {"business_date": ["2026-10-01", "2026-10-31"]}},
        {"query": {"view": "movement", "warning": "high"}},
        {"unexpected": True},
    ],
)
def test_invalid_export_options(test_client, auth_headers, options):
    """文件名、字段、格式和跨视图条件必须由服务端校验。"""
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json={"fields": ["goods_code"], **options})
    assert response.status_code == 422, response.text


def test_export_text_formula_guard_and_duplicate_fields(test_client, auth_headers, stock_report_data):
    """档案公式文本输出为字面值，字段去重不乱序，负金额不加公式防护前缀。"""
    import csv
    from io import StringIO

    async def rename():
        async with async_db_session() as db, db.begin():
            row = await db.get(GoodsModel, stock_report_data["goods"][0])
            row.goods_name = "=2+2"

    test_client.portal.call(rename)
    body = {
        "query": {"view": "movement", "warehouse_id": stock_report_data["warehouses"][1]},
        "fields": ["goods_name", "net_amount", "goods_name"],
        "format": "csv",
    }
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    values = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
    assert values == [["商品名称", "净变动金额"], ["'=2+2", "-5.00"]]
    body["format"] = "xlsx"
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.cell(2, 1).value == "'=2+2" and sheet.cell(2, 1).data_type == "s"
    assert sheet.cell(2, 2).value == "-5.00"


def test_export_over_limit_is_explicit(test_client):
    """超过框架十万行限制明确拒绝，不生成只含部分数据的文件。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.exceptions import CustomException
    from app.modules.inventory.report.stock.schema import StockReportExport
    from app.modules.inventory.report.stock.service import StockReportService

    async def run():
        async with async_db_session() as db:
            auth = AuthSchema(user=CoreUserSchema(id=1, username="super", name="super", dept_id=None, is_superuser=True))
            service = StockReportService(auth, db)
            with patch.object(service.crud, "rows", new=AsyncMock(return_value=([], 100001))):
                with pytest.raises(CustomException, match="10万行"):
                    await service.export(StockReportExport(fields=["goods_code"]))

    test_client.portal.call(run)


def test_stock_report_permission_is_independent_and_readonly(test_client, auth_headers, stock_report_data):
    """普通角色可独立查询企业库存报表，但不继承库存明细、导出或写入权限。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.dependencies import get_current_user

    async def load_user():
        async with async_db_session() as db:
            user = await db.scalar(select(UserModel).where(UserModel.username == "user"))
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    auth = AuthSchema(user=test_client.portal.call(load_user), permissions=["module_inventory:stock_report:query"])
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        assert data(test_client, {}, stock_report_data, "/summary")["stock_amount"] == "103.00"
        assert data(test_client, {}, stock_report_data, "/summary", {"view": "movement"})["net_amount"] == "70.01"
        assert test_client.get("/inventory/stock/balance/list").status_code == 403
        assert test_client.get("/inventory/stock/log/list").status_code == 403
        assert test_client.post(PREFIX + "/export", json={"fields": ["goods_code"]}).status_code == 403
        auth.permissions = ["module_inventory:stock_balance:query", "module_inventory:stock_log:query"]
        assert test_client.get(PREFIX + "/list").status_code == 403
        assert test_client.get(PREFIX + "/summary").status_code == 403
        auth.permissions = ["module_inventory:stock_report:export"]
        response = test_client.post(PREFIX + "/export", json={"query": {"warehouse_name": stock_report_data["key"]}, "fields": ["goods_code"]})
        assert response.status_code == 200, response.text
        assert test_client.get(PREFIX + "/list").status_code == 403
        for endpoint in ["create", "update/1", "delete", "confirm/1", "void/1"]:
            assert test_client.post(PREFIX + "/" + endpoint, json={}).status_code in (404, 405)
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous


def test_all_export_views_handle_empty_results(test_client, auth_headers, stock_report_data):
    """零结果仍输出合法文件和所选表头，不误报导出失败。"""
    for view, field in [("balance", "stock_amount"), ("movement", "net_amount")]:
        response = test_client.post(
            PREFIX + "/export",
            headers=auth_headers,
            json={
                "query": {"view": view, "warehouse_name": stock_report_data["key"], "goods_code": "unmatched"},
                "fields": [field],
                "format": "csv",
            },
        )
        assert response.status_code == 200, response.text
        assert len(response.content.decode("utf-8-sig").splitlines()) == 1


def test_aggregate_amount_exceeds_single_row_limit_without_validation_cap(test_client, auth_headers, stock_report_data):
    """汇总不套用单行金额上限，临时 SQLite 使用可精确表示的大整数验证。"""

    async def change():
        async with async_db_session() as db, db.begin():
            rows = (
                await db.scalars(
                    select(StockBalanceModel).where(
                        StockBalanceModel.warehouse_id.in_(stock_report_data["warehouses"]),
                        StockBalanceModel.goods_id == stock_report_data["goods"][1],
                    )
                )
            ).all()
            for row in rows:
                row.stock_amount = Decimal("9000000000000000.00")

    test_client.portal.call(change)
    summary = data(test_client, auth_headers, stock_report_data, "/summary", {"goods_code": stock_report_data["codes"][1]})
    assert summary["stock_amount"] == "18000000000000000.00"
