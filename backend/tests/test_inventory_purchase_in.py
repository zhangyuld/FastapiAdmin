import csv
import json
import re
from collections.abc import Generator
from copy import deepcopy
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any
from unittest.mock import AsyncMock, patch
from urllib.parse import unquote
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, func, select, update

from app.core.base_schema import AuthSchema, CoreUserSchema
from app.core.database import async_db_session
from app.core.dependencies import get_current_user
from app.core.exceptions import CustomException
from app.modules.inventory.basic.model import GoodsModel
from app.modules.inventory.common.business_log import BusinessLogService
from app.modules.inventory.common.model import BusinessLogModel
from app.modules.inventory.finance.model import FundAllocationModel, FundTransactionModel, PayableModel
from app.modules.inventory.finance.payable import PayableService
from app.modules.inventory.purchase.inbound import crud as crud_module
from app.modules.inventory.purchase.inbound import service as service_module
from app.modules.inventory.purchase.model import PurchaseInItemModel, PurchaseInModel, PurchaseReturnItemModel, PurchaseReturnModel
from app.modules.inventory.stock.model import StockBalanceModel, StockLogModel
from app.modules.inventory.stock.posting import StockPostingService
from app.modules.system.user.model import UserModel

BASE = "/inventory/purchase/inbound"


@pytest.fixture(scope="session", autouse=True)
def disable_purchase_test_audit_writes() -> Generator[None, None, None]:
    """SQLite 测试中隔离另会话的审计写入，避免阻塞尚未提交的业务事务。"""
    with (
        patch("app.core.router_class._write_operation_log_async", new=AsyncMock()),
        patch("app.modules.system.auth.service._write_login_log", new=AsyncMock()),
    ):
        yield


@pytest.fixture
def purchase_body(test_client: TestClient, auth_headers: dict[str, str]) -> dict[str, Any]:
    """创建互不冲突的启用档案，并生成两条采购明细入参。"""
    code = uuid4().hex[:12]
    category = test_client.post("/inventory/basic/category/create", headers=auth_headers, json={"name": "采购分类" + code, "parent_id": None, "sort": 1, "status": 0})
    assert category.status_code == 200, category.text
    supplier = test_client.post(
        "/inventory/basic/supplier/create",
        headers=auth_headers,
        json={
            "supplier_name": "采购供应商" + code,
            "supplier_code": "SUP" + code,
            "contact": "测试联系人",
            "phone": "13800138000",
            "status": 0,
        },
    )
    assert supplier.status_code == 200, supplier.text
    warehouse = test_client.post("/inventory/basic/warehouse/create", headers=auth_headers, json={"warehouse_name": "采购仓库" + code, "warehouse_code": "WH" + code, "status": 0})
    assert warehouse.status_code == 200, warehouse.text
    goods_ids = []
    for index in range(2):
        response = test_client.post(
            "/inventory/basic/goods/create",
            headers=auth_headers,
            json={
                "goods_code": f"GOODS{code}{index}",
                "goods_name": f"采购商品{index}",
                "category_id": category.json()["data"]["id"],
                "spec": "规格A",
                "unit": "件",
                "reference_cost_price": "1",
                "sale_price": "2",
                "stock_min": "0",
                "stock_max": "0",
                "status": 0,
            },
        )
        assert response.status_code == 200, response.text
        goods_ids.append(response.json()["data"]["id"])
    return {
        "supplier_id": supplier.json()["data"]["id"],
        "warehouse_id": warehouse.json()["data"]["id"],
        "business_date": "2026-10-12",
        "remark": "  采购草稿  ",
        "items": [{"goods_id": id, "quantity": "1.000", "unit_price": "0.0050"} for id in goods_ids],
    }


def _create(client: TestClient, headers: dict[str, str], body: dict[str, Any]) -> dict[str, Any]:
    """创建草稿并要求接口正确返回明细，而非只返回主单。"""
    response = client.post(BASE + "/create", headers=headers, json=body)
    assert response.status_code == 201, response.text
    return response.json()["data"]


def test_purchase_draft_lifecycle(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """草稿完成创建、筛选、更新、版本冲突及主从表软删除闭环。"""
    document = _create(test_client, auth_headers, purchase_body)
    id = document["id"]
    assert document["total_amount"] == "0.02"
    assert document["status"] == 0 and document["version"] == 0
    assert document["remark"] == "采购草稿"
    assert document["supplier_name"].startswith("采购供应商")
    assert document["warehouse_name"].startswith("采购仓库")
    assert document["created_by"]["name"]
    assert [item["line_no"] for item in document["items"]] == [1, 2]
    assert [item["line_amount"] for item in document["items"]] == ["0.01", "0.01"]
    assert document["items"][0]["quantity"] == "1.000"
    assert document["items"][0]["unit_price"] == "0.0050"
    assert document["items"][0]["goods_code"].startswith("GOODS")

    listed = test_client.get(
        BASE + "/list",
        headers=auth_headers,
        params=[
            ("supplier_id", purchase_body["supplier_id"]),
            ("warehouse_id", purchase_body["warehouse_id"]),
            ("order_no", document["order_no"]),
            ("business_date", "2026-10-12"),
            ("business_date", "2026-10-12"),
            ("created_id", document["created_id"]),
            ("status", "0"),
            ("page_no", 1),
            ("page_size", 1),
        ],
    )
    assert listed.status_code == 200, listed.text
    assert listed.json()["data"]["total"] == 1
    assert "items" not in listed.json()["data"]["items"][0]
    excluded = test_client.get(BASE + "/list", headers=auth_headers, params=[("order_no", document["order_no"]), ("business_date", "2026-10-13"), ("business_date", "2026-10-14")])
    assert excluded.json()["data"]["total"] == 0

    body = deepcopy(purchase_body)
    body.update(version=0, remark="")
    body["items"] = [{"goods_id": body["items"][1]["goods_id"], "quantity": "2.500", "unit_price": "1.2345"}]
    changed = test_client.post(BASE + f"/update/{id}", headers=auth_headers, json=body)
    assert changed.status_code == 200, changed.text
    current = changed.json()["data"]
    assert current["order_no"] == document["order_no"]
    assert current["version"] == 1 and current["remark"] is None
    assert current["total_amount"] == "3.09"
    assert len(current["items"]) == 1 and current["items"][0]["line_no"] == 1
    stale = test_client.post(BASE + f"/update/{id}", headers=auth_headers, json=body)
    assert stale.status_code == 409, stale.text
    assert test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [id], "versions": {id: 0}}).status_code == 409
    deleted = test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [id, id], "versions": {id: 1}})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["data"] == [id]
    assert test_client.get(BASE + f"/detail/{id}", headers=auth_headers).status_code == 404

    async def inspect_deletion() -> tuple[bool, bool, int, int]:
        """直接检查主单与明细软删除，以及草稿未产生库存、应付。"""
        async with async_db_session() as db:
            main = await db.get(PurchaseInModel, id)
            items = (await db.execute(select(PurchaseInItemModel).where(PurchaseInItemModel.purchase_in_id == id))).scalars().all()
            stock = await db.scalar(select(func.count()).select_from(StockBalanceModel).where(StockBalanceModel.warehouse_id == purchase_body["warehouse_id"]))
            payable = await db.scalar(select(func.count()).select_from(PayableModel).where(PayableModel.supplier_id == purchase_body["supplier_id"]))
            return main.is_deleted, all(item.is_deleted for item in items), stock, payable

    assert test_client.portal.call(inspect_deletion) == (True, True, 0, 0)


@pytest.mark.parametrize("kind", ["supplier", "warehouse", "goods"])
@pytest.mark.parametrize("deleted", [False, True])
def test_purchase_rejects_inactive_archives(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], kind: str, deleted: bool) -> None:
    """拒绝停用及已删除的供应商、仓库和商品，失败不会产生主单。"""
    id = purchase_body["items"][0]["goods_id"] if kind == "goods" else purchase_body[kind + "_id"]
    path = f"/inventory/basic/{kind}/" + ("delete" if deleted else "status/batch")
    payload = {"ids": [id]} if deleted else {"ids": [id], "status": 1}
    if kind == "goods" and not deleted:

        async def disable_goods() -> None:
            """现有商品模块没有状态接口，测试直接预置停用状态。"""
            async with async_db_session() as db, db.begin():
                await db.execute(update(GoodsModel).where(GoodsModel.id == id).values(status=1))

        test_client.portal.call(disable_goods)
    else:
        response = test_client.post(path, headers=auth_headers, json=payload)
        assert response.status_code == 200, response.text
    result = test_client.post(BASE + "/create", headers=auth_headers, json=purchase_body)
    assert result.status_code == (404 if deleted else 400), result.text
    listed = test_client.get(BASE + "/list", headers=auth_headers, params={"supplier_id": purchase_body["supplier_id"]})
    assert listed.json()["data"]["total"] == 0


def test_purchase_export_selected_columns(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """全量导出保留筛选、字段顺序、显示状态和自定义或默认文件名。"""
    document = _create(test_client, auth_headers, purchase_body)
    exported = test_client.post(
        BASE + "/export",
        headers=auth_headers,
        json={
            "query": {"order_no": document["order_no"]},
            "fields": ["total_amount", "order_no", "status", "created_name"],
            "format": "xlsx",
            "sheet_name": "采购测试",
            "filename": "采购记录.xlsx",
        },
    )
    assert exported.status_code == 200, exported.text
    assert "filename=purchase_in.xlsx; filename*=" in exported.headers["content-disposition"]
    assert "采购记录.xlsx" in unquote(exported.headers["content-disposition"])
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    rows = list(workbook["采购测试"].iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == ("总金额", "单据编号", "状态", "创建人")
    assert len(rows) == 2
    assert rows[1][:3] == ("0.02", document["order_no"], "草稿")
    assert rows[1][3] == document["created_by"]["name"]

    csv_file = test_client.post(
        BASE + "/export",
        headers=auth_headers,
        json={
            "query": {"order_no": document["order_no"]},
            "fields": ["order_no", "supplier_name"],
            "format": "csv",
            "filename": "",
        },
    )
    assert csv_file.status_code == 200, csv_file.text
    assert re.search(r"采购入库列表_\d{8}_\d{6}\.csv", unquote(csv_file.headers["content-disposition"]))
    csv_rows = list(csv.reader(StringIO(csv_file.content.decode("utf-8-sig"))))
    assert csv_rows == [["单据编号", "供应商"], [document["order_no"], document["supplier_name"]]]


def test_purchase_rejects_amount_overflow_and_bad_sort(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """明细精度虽合法，乘积或合计超数据库范围仍须拒绝。"""
    purchase_body["items"] = [{"goods_id": purchase_body["items"][0]["goods_id"], "quantity": "999999999999999.999", "unit_price": "99999999999999.9999"}]
    result = test_client.post(BASE + "/create", headers=auth_headers, json=purchase_body)
    assert result.status_code == 400, result.text
    assert "采购金额超过允许范围" in result.json()["msg"]
    invalid = test_client.get(BASE + "/list", headers=auth_headers, params={"order_by": json.dumps([{"supplier_name": "desc"}])})
    assert invalid.status_code == 422, invalid.text


def test_purchase_total_amount_overflow(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """各行金额未超限但合计超限时，仍应在写主从表前拒绝。"""
    for item in purchase_body["items"]:
        item.update(quantity="999999999999999.999", unit_price="10.0000")
    response = test_client.post(BASE + "/create", headers=auth_headers, json=purchase_body)
    assert response.status_code == 400, response.text
    listed = test_client.get(BASE + "/list", headers=auth_headers, params={"supplier_id": purchase_body["supplier_id"]})
    assert listed.json()["data"]["total"] == 0


def test_purchase_goods_snapshot_preserved(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """商品档案改名后，已保存的明细仍展示原有快照。"""
    document = _create(test_client, auth_headers, purchase_body)

    async def rename_goods() -> None:
        """直接变更档案名称，验证单据没有读取实时商品名称。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(GoodsModel).where(GoodsModel.id == purchase_body["items"][0]["goods_id"]).values(goods_name="更名后的商品"))

    test_client.portal.call(rename_goods)
    current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
    assert current["items"][0]["goods_name"] == document["items"][0]["goods_name"]


def test_purchase_crud_atomic_version_guard(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """越过业务层的早期版本校验，SQL 本身仍拒绝旧版本修改和删除。"""
    document = _create(test_client, auth_headers, purchase_body)
    changed = test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json={**purchase_body, "version": 0})
    assert changed.status_code == 200, changed.text

    async def submit_stale_version() -> tuple[bool, bool]:
        """使用独立事务模拟已持有旧版本的请求进行条件写入。"""
        async with async_db_session() as db, db.begin():
            auth = AuthSchema(user=CoreUserSchema(id=document["created_id"], is_superuser=True))
            crud = crud_module.PurchaseInCRUD(auth, db)
            updated = await crud.update_draft(document["id"], 0, {"remark": "不应写入"}, [])
            deleted = await crud.delete_draft(document["id"], 0)
            return updated, deleted

    assert test_client.portal.call(submit_stale_version) == (False, False)
    current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
    assert current["version"] == 1 and current["remark"] == "采购草稿"
    assert len(current["items"]) == 2


@pytest.mark.parametrize("state", [1, 2])
def test_purchase_non_draft_readonly(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], state: int) -> None:
    """模拟后续生效或作废状态，现阶段接口不能修改或删除这些单据。"""
    document = _create(test_client, auth_headers, purchase_body)

    async def set_status() -> None:
        """仅测试预置状态，不调用尚未实现的确认或作废功能。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(PurchaseInModel).where(PurchaseInModel.id == document["id"]).values(status=state))

    test_client.portal.call(set_status)
    updated = test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json={**purchase_body, "version": 0})
    deleted = test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [document["id"]], "versions": {document["id"]: 0}})
    assert updated.status_code == 409 and deleted.status_code == 409
    assert test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]["status"] == state


@pytest.mark.parametrize("operation", ["create", "update"])
def test_purchase_main_and_items_rollback(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], monkeypatch: pytest.MonkeyPatch, operation: str) -> None:
    """明细写入后出错，创建和替换明细都必须连同主单整体回滚。"""
    document = _create(test_client, auth_headers, purchase_body) if operation == "update" else None
    order_no = "ROLLBACK" + uuid4().hex
    monkeypatch.setattr(service_module.PurchaseInService, "_generate_order_no", staticmethod(lambda: order_no))
    original = crud_module.PurchaseInCRUD._add_items

    async def fail_after_items(self: Any, id: int, items: list[Any]) -> None:
        """模拟明细已经写入但事务提交前发生业务异常。"""
        await original(self, id, items)
        raise CustomException(msg="模拟明细写入失败", status_code=409)

    monkeypatch.setattr(crud_module.PurchaseInCRUD, "_add_items", fail_after_items)
    body = deepcopy(purchase_body)
    body["remark"] = "不应保存"
    if document:
        body.update(version=0)
        body["items"] = body["items"][:1]
        response = test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json=body)
    else:
        response = test_client.post(BASE + "/create", headers=auth_headers, json=body)
    assert response.status_code == 409, response.text
    if document:
        current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
        assert current["version"] == 0 and current["remark"] == document["remark"]
        assert [item["id"] for item in current["items"]] == [item["id"] for item in document["items"]]
    else:
        listed = test_client.get(BASE + "/list", headers=auth_headers, params={"order_no": order_no})
        assert listed.json()["data"]["total"] == 0


def test_purchase_batch_delete_rollback(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
    """批量删除第一张已写入后第二张发生版本冲突，整批仍须回滚。"""
    documents = [_create(test_client, auth_headers, purchase_body) for _ in range(2)]
    original = crud_module.PurchaseInCRUD.delete_draft
    calls = 0

    async def conflict_on_second(self: Any, id: int, version: int) -> bool:
        """模拟第二张草稿的原子版本校验失败。"""
        nonlocal calls
        calls += 1
        return False if calls == 2 else await original(self, id, version)

    monkeypatch.setattr(crud_module.PurchaseInCRUD, "delete_draft", conflict_on_second)
    result = test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [item["id"] for item in documents], "versions": {item["id"]: 0 for item in documents}})
    assert result.status_code == 409, result.text
    for document in documents:
        current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers)
        assert current.status_code == 200, current.text
        assert current.json()["data"]["version"] == 0


def test_purchase_data_scope_and_button_permission(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """未登录不可访问；真实普通用户无权查询、修改、删除或导出他人的单据。"""
    document = _create(test_client, auth_headers, purchase_body)
    assert test_client.get(BASE + "/list").status_code == 401

    async def load_user() -> CoreUserSchema:
        """加载数据范围为仅本人的种子用户。"""
        async with async_db_session() as db:
            user = (await db.execute(select(UserModel).where(UserModel.username == "user"))).scalar_one()
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    auth = AuthSchema(
        user=test_client.portal.call(load_user), permissions=[f"module_inventory:purchase_in:{action}" for action in ["query", "detail", "create", "update", "delete", "export", "confirm", "void"]]
    )
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        listed = test_client.get(BASE + "/list", params={"order_no": document["order_no"]})
        assert listed.status_code == 200 and listed.json()["data"]["total"] == 0
        assert test_client.get(BASE + f"/detail/{document['id']}").status_code == 404
        assert test_client.get(BASE + f"/returnable/{document['id']}").status_code == 404
        assert test_client.post(BASE + f"/confirm/{document['id']}", json={"version": 0}).status_code == 404
        assert test_client.post(BASE + f"/void/{document['id']}", json={"version": 0, "void_reason": "测试"}).status_code == 404
        assert test_client.post(BASE + f"/update/{document['id']}", json={**purchase_body, "version": 0}).status_code == 404
        assert test_client.post(BASE + "/delete", json={"ids": [document["id"]], "versions": {document["id"]: 0}}).status_code == 404
        exported = test_client.post(BASE + "/export", json={"query": {"order_no": document["order_no"]}, "fields": ["order_no"], "format": "csv"})
        assert list(csv.reader(StringIO(exported.content.decode("utf-8-sig")))) == [["单据编号"]]
        own = _create(test_client, {}, purchase_body)
        assert own["created_id"] == auth.user.id
        assert test_client.get(BASE + f"/detail/{own['id']}").status_code == 200
        auth.permissions = []
        assert test_client.get(BASE + "/list").status_code == 403
        assert test_client.get(BASE + f"/returnable/{own['id']}").status_code == 403
        assert test_client.post(BASE + f"/confirm/{own['id']}", json={"version": 0}).status_code == 403
        assert test_client.post(BASE + f"/void/{own['id']}", json={"version": 0, "void_reason": "测试"}).status_code == 403
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous


def _confirm(client: TestClient, headers: dict[str, str], document: dict[str, Any]) -> dict[str, Any]:
    """确认单据并要求接口返回最新状态和明细。"""
    response = client.post(BASE + f"/confirm/{document['id']}", headers=headers, json={"version": document["version"]})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _accounting(client: TestClient, body: dict[str, Any]) -> dict[str, Any]:
    """读取当前测试仓库的完整记账快照，用于入账和回滚断言。"""

    async def inspect() -> dict[str, Any]:
        """同一测试事件循环读取库存、流水、应付和业务日志。"""
        async with async_db_session() as db:
            balances = (await db.execute(select(StockBalanceModel).where(StockBalanceModel.warehouse_id == body["warehouse_id"]).order_by(StockBalanceModel.goods_id))).scalars().all()
            logs = (await db.execute(select(StockLogModel).where(StockLogModel.warehouse_id == body["warehouse_id"]).order_by(StockLogModel.id))).scalars().all()
            bills = (await db.execute(select(PayableModel).where(PayableModel.supplier_id == body["supplier_id"]).order_by(PayableModel.id))).scalars().all()
            document_ids = select(PurchaseInModel.id).where(PurchaseInModel.warehouse_id == body["warehouse_id"])
            business_logs = (
                (await db.execute(select(BusinessLogModel).where(BusinessLogModel.business_type == "purchase_in", BusinessLogModel.business_id.in_(document_ids)).order_by(BusinessLogModel.id)))
                .scalars()
                .all()
            )
            return {
                "balances": [(row.quantity, row.stock_amount, row.avg_cost, row.version) for row in balances],
                "logs": [(row.id, row.business_item_id, row.operation_type, row.change_quantity, row.cost_amount, row.reversal_log_id, row.is_deleted) for row in logs],
                "bills": [(row.id, row.operation_type, row.amount, row.allocated_amount, row.unallocated_amount, row.reversal_id, row.status, row.is_deleted) for row in bills],
                "audit": [(row.action, row.before_status, row.after_status, row.content) for row in business_logs],
            }

    return client.portal.call(inspect)


def test_purchase_confirmation_and_reversal(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """确认、重复请求、版本冲突及完整冲销保留原流水并生成关联反向流水。"""
    document = _create(test_client, auth_headers, purchase_body)
    assert test_client.get(BASE + f"/returnable/{document['id']}", headers=auth_headers).status_code == 409
    assert test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 0, "void_reason": "测试"}).status_code == 409
    confirmed = _confirm(test_client, auth_headers, document)
    assert confirmed["status"] == 1 and confirmed["version"] == 1
    assert confirmed["confirmed_id"] and confirmed["confirmed_time"] and confirmed["confirmed_name"]
    assert [
        {key: value for key, value in item.items() if key != "returnable_quantity"}
        for item in confirmed["items"]
    ] == [
        {key: value for key, value in item.items() if key != "returnable_quantity"}
        for item in document["items"]
    ]
    assert all(item["returnable_quantity"] == item["quantity"] for item in confirmed["items"])
    before = _accounting(test_client, purchase_body)
    assert before["balances"] == [(Decimal("1.000"), Decimal("0.01"), Decimal("0.0100"), 1)] * 2
    assert len(before["logs"]) == 2 and before["bills"][0][2:5] == (Decimal("0.02"), Decimal("0.00"), Decimal("0.02"))
    assert before["audit"] == [("confirm", 0, 1, None)]
    assert test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0}).status_code == 409
    assert test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 0, "void_reason": "测试"}).status_code == 409
    assert _accounting(test_client, purchase_body) == before
    source = test_client.get(BASE + f"/returnable/{document['id']}", headers=auth_headers)
    assert source.status_code == 200, source.text
    assert [item["returnable_quantity"] for item in source.json()["data"]["items"]] == ["1.000", "1.000"]
    result = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "  录入错误  "})
    assert result.status_code == 200, result.text
    current = result.json()["data"]
    assert current["status"] == 2 and current["version"] == 2 and current["void_reason"] == "录入错误"
    assert current["voided_id"] and current["voided_time"] and current["voided_name"]
    after = _accounting(test_client, purchase_body)
    assert after["balances"] == [(Decimal("0"), Decimal("0"), Decimal("0"), 2)] * 2
    assert after["logs"][:2] == before["logs"]
    original_map = {row[1]: row for row in before["logs"]}
    for row in after["logs"][2:]:
        assert row[2] == "void" and row[3:5] == (-original_map[row[1]][3], -original_map[row[1]][4])
        assert row[5] == original_map[row[1]][0] and not row[6]
    assert after["bills"][0] == before["bills"][0]
    assert after["bills"][1][2] == Decimal("-0.02") and after["bills"][1][5] == after["bills"][0][0]
    assert after["audit"] == [("confirm", 0, 1, None), ("void", 1, 2, "录入错误")]
    assert test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "重复"}).status_code == 409
    assert test_client.get(BASE + f"/returnable/{document['id']}", headers=auth_headers).status_code == 409
    assert _accounting(test_client, purchase_body) == after


def test_purchase_moving_average_and_original_cost_reversal(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """连续入库按金额加权，冲销较早入库时仍按原成本而非当前平均成本。"""
    goods_id = purchase_body["items"][0]["goods_id"]
    purchase_body["items"] = [{"goods_id": goods_id, "quantity": "10", "unit_price": "2"}]
    first = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    purchase_body["items"] = [{"goods_id": goods_id, "quantity": "5", "unit_price": "4"}]
    _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    assert _accounting(test_client, purchase_body)["balances"] == [(Decimal("15"), Decimal("40"), Decimal("2.6667"), 2)]
    result = test_client.post(BASE + f"/void/{first['id']}", headers=auth_headers, json={"version": 1, "void_reason": "冲销第一笔"})
    assert result.status_code == 200, result.text
    assert _accounting(test_client, purchase_body)["balances"] == [(Decimal("5"), Decimal("20"), Decimal("4"), 3)]


@pytest.mark.parametrize("operation", ["confirm", "void"])
@pytest.mark.parametrize("stage", ["stock", "payable", "audit"])
def test_purchase_posting_rollback(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], monkeypatch: pytest.MonkeyPatch, operation: str, stage: str) -> None:
    """库存、应付或日志已写入后失败，余额、流水、状态和版本必须全部回滚。"""
    document = _create(test_client, auth_headers, purchase_body)
    if operation == "void":
        document = _confirm(test_client, auth_headers, document)
    before = _accounting(test_client, purchase_body)
    cls, method = {
        "stock": (StockPostingService, "post"),
        "payable": (PayableService, "reverse" if operation == "void" else "create"),
        "audit": (BusinessLogService, "record"),
    }[stage]
    original = getattr(cls, method)

    async def fail_after_write(self: Any, *args: Any, **kwargs: Any) -> None:
        """注入写入后的业务失败以检查整个请求事务。"""
        await original(self, *args, **kwargs)
        raise CustomException(msg="模拟记账失败", status_code=409)

    monkeypatch.setattr(cls, method, fail_after_write)
    body = {"version": document["version"]}
    if operation == "void":
        body["void_reason"] = "测试冲销"
    response = test_client.post(BASE + f"/{operation}/{document['id']}", headers=auth_headers, json=body)
    assert response.status_code == 409, response.text
    current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
    assert current["status"] == document["status"] and current["version"] == document["version"]
    assert current["confirmed_time"] == document["confirmed_time"] and current["voided_time"] == document["voided_time"]
    assert _accounting(test_client, purchase_body) == before


@pytest.mark.parametrize("quantity,amount", [("0.500", "0.01"), ("1", "0"), ("1", "0.02")])
def test_purchase_void_rejects_insufficient_or_residual_stock(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], quantity: str, amount: str) -> None:
    """数量不足、金额不足、清空数量后残留金额均不能强行冲销。"""
    document = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))

    async def consume_stock() -> None:
        """模拟其他库存业务已经影响原余额。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(StockBalanceModel).where(StockBalanceModel.warehouse_id == purchase_body["warehouse_id"]).values(quantity=Decimal(quantity), stock_amount=Decimal(amount)))

    test_client.portal.call(consume_stock)
    before = _accounting(test_client, purchase_body)
    result = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "不能冲销"})
    assert result.status_code == 409, result.text
    assert _accounting(test_client, purchase_body) == before
    assert test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]["status"] == 1


@pytest.mark.parametrize("allocation", [False, True])
def test_purchase_void_rejects_settlement(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], allocation: bool) -> None:
    """已核销金额或未撤销的核销记录任一存在，都必须阻止作废。"""
    document = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))

    async def settle() -> None:
        """独立预置金额或核销记录，避免只检查其中一个字段。"""
        async with async_db_session() as db, db.begin():
            bill = (await db.execute(select(PayableModel).where(PayableModel.business_type == "purchase_in", PayableModel.business_id == document["id"]))).scalar_one()
            if allocation:
                fund = FundTransactionModel(
                    transaction_no=uuid4().hex,
                    transaction_type="payment",
                    party_type="supplier",
                    party_id=purchase_body["supplier_id"],
                    amount=Decimal("0.01"),
                    payment_method="cash",
                    business_date=date(2026, 10, 12),
                    status=1,
                )
                db.add(fund)
                await db.flush()
                db.add(FundAllocationModel(fund_transaction_id=fund.id, bill_type="payable", bill_id=bill.id, allocated_amount=Decimal("0.01")))
            else:
                bill.allocated_amount = Decimal("0.01")
                bill.unallocated_amount = Decimal("0.01")
                bill.status = 1

    test_client.portal.call(settle)
    before = _accounting(test_client, purchase_body)
    response = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "核销测试"})
    assert response.status_code == 409, response.text
    assert "核销" in response.json()["msg"] and _accounting(test_client, purchase_body) == before


def _seed_return(client: TestClient, document: dict[str, Any], state: int, quantity: str, *, deleted: bool = False) -> int:
    """预置后续退货单及来源明细，不将尚未开发的退货接口纳入本阶段。"""

    async def create() -> int:
        """复制来源快照建立退货测试数据。"""
        async with async_db_session() as db, db.begin():
            main = PurchaseReturnModel(
                order_no="RETURN" + uuid4().hex,
                supplier_id=document["supplier_id"],
                warehouse_id=document["warehouse_id"],
                source_purchase_in_id=document["id"],
                business_date=date(2026, 10, 12),
                total_amount=Decimal("0"),
                return_reason="测试退货",
                status=state,
                is_deleted=deleted,
                created_id=document["created_id"],
            )
            db.add(main)
            await db.flush()
            item = document["items"][0]
            db.add(
                PurchaseReturnItemModel(
                    purchase_return_id=main.id,
                    source_purchase_in_item_id=item["id"],
                    line_no=1,
                    goods_id=item["goods_id"],
                    goods_code=item["goods_code"],
                    goods_name=item["goods_name"],
                    spec=item["spec"],
                    unit=item["unit"],
                    quantity=Decimal(quantity),
                    unit_price=Decimal(item["unit_price"]),
                    unit_cost=Decimal(item["unit_price"]),
                    line_amount=Decimal("0"),
                    cost_amount=Decimal("0"),
                )
            )
            return main.id

    return client.portal.call(create)


def test_purchase_returnable_rules_and_excluded_draft(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """汇总多笔有效退货，忽略草稿、作废和删除记录，验证编辑排除参数归属。"""
    document = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    draft_id = _seed_return(test_client, document, 0, "0.800")
    confirmed_id = _seed_return(test_client, document, 1, "0.200")
    _seed_return(test_client, document, 1, "0.100")
    _seed_return(test_client, document, 2, "0.900")
    _seed_return(test_client, document, 1, "0.900", deleted=True)
    result = test_client.get(BASE + f"/returnable/{document['id']}", headers=auth_headers, params={"exclude_return_id": draft_id})
    assert result.status_code == 200, result.text
    assert [item["returnable_quantity"] for item in result.json()["data"]["items"]] == ["0.700", "1.000"]
    assert test_client.get(BASE + f"/returnable/{document['id']}", headers=auth_headers, params={"exclude_return_id": confirmed_id}).status_code == 404
    other = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    assert test_client.get(BASE + f"/returnable/{other['id']}", headers=auth_headers, params={"exclude_return_id": draft_id}).status_code == 404
    assert test_client.get(BASE + f"/returnable/{other['id']}", headers=auth_headers, params={"exclude_return_id": 999999}).status_code == 404
    assert test_client.get(BASE + f"/returnable/{other['id']}", headers=auth_headers, params={"exclude_return_id": 0}).status_code == 422


@pytest.mark.parametrize("state,deleted,expected", [(0, False, 409), (1, False, 409), (2, False, 200), (0, True, 200)])
def test_purchase_void_linked_return_guard(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], state: int, deleted: bool, expected: int) -> None:
    """任何未作废且未删除的来源退货阻止入库作废，避免留下无效来源草稿。"""
    document = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    _seed_return(test_client, document, state, "0.100", deleted=deleted)
    result = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "关联退货检查"})
    assert result.status_code == expected, result.text


def test_purchase_zero_price_and_confirm_archive_validation(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """确认时再次校验档案；零单价入库正常增加数量且生成已结清零应付。"""
    purchase_body["items"][0]["unit_price"] = "0"
    purchase_body["items"] = purchase_body["items"][:1]
    document = _create(test_client, auth_headers, purchase_body)

    async def set_goods_status(value: int) -> None:
        """模拟草稿录入后商品被停用或恢复。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(GoodsModel).where(GoodsModel.id == purchase_body["items"][0]["goods_id"]).values(status=value))

    test_client.portal.call(set_goods_status, 1)
    result = test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0})
    assert result.status_code == 400, result.text
    assert _accounting(test_client, purchase_body)["balances"] == []
    test_client.portal.call(set_goods_status, 0)
    _confirm(test_client, auth_headers, document)
    values = _accounting(test_client, purchase_body)
    assert values["balances"] == [(Decimal("1"), Decimal("0"), Decimal("0"), 1)]
    assert values["bills"][0][2:5] == (Decimal("0"), Decimal("0"), Decimal("0")) and values["bills"][0][6] == 2


@pytest.mark.parametrize("field,value", [("quantity", "999999999999999.999"), ("stock_amount", "9999999999999999.99")])
def test_purchase_confirmation_rejects_balance_overflow(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], field: str, value: str) -> None:
    """累计余额超出数据库精度时拒绝确认，不能只检查本次采购金额。"""
    _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))

    async def fill_balance() -> None:
        """预置接近上限的有效余额。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(StockBalanceModel).where(StockBalanceModel.warehouse_id == purchase_body["warehouse_id"]).values(**{field: Decimal(value)}))

    test_client.portal.call(fill_balance)
    document = _create(test_client, auth_headers, purchase_body)
    before = _accounting(test_client, purchase_body)
    result = test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0})
    assert result.status_code == 400, result.text
    assert _accounting(test_client, purchase_body) == before


@pytest.mark.parametrize("missing", ["log", "payable"])
def test_purchase_void_rejects_missing_original_posting(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], missing: str) -> None:
    """历史入账不完整不能凭单据直接冲销或补造原流水。"""
    document = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))

    async def remove_original() -> None:
        """模拟被错误删除的原入账数据。"""
        model = StockLogModel if missing == "log" else PayableModel
        async with async_db_session() as db, db.begin():
            await db.execute(delete(model).where(model.business_type == "purchase_in", model.business_id == document["id"]))

    test_client.portal.call(remove_original)
    before = _accounting(test_client, purchase_body)
    response = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "原记录缺失"})
    assert response.status_code == 409, response.text
    assert _accounting(test_client, purchase_body) == before


def test_purchase_state_transition_rejects_stale_orm(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """即使业务层持有旧 ORM 版本，数据库状态更新仍原子阻止覆盖。"""
    document = _create(test_client, auth_headers, purchase_body)

    async def submit_stale() -> None:
        """模拟读取后另一次更新先行完成，校验条件更新而非仅内存判断。"""
        async with async_db_session() as db, db.begin():
            user = await db.get(UserModel, document["created_id"])
            auth = AuthSchema(user=CoreUserSchema(id=user.id, username=user.username, name=user.name, is_superuser=True))
            crud = crud_module.PurchaseInCRUD(auth, db)
            stale = await crud.get_document(document["id"])
            await db.execute(update(PurchaseInModel).where(PurchaseInModel.id == stale.id).values(version=1).execution_options(synchronize_session=False))
            with pytest.raises(CustomException):
                await crud.transition(stale, service_module.DocumentStatus.CONFIRMED, datetime.now(UTC))
            current = await crud.get_document(document["id"])
            assert current.version == 1 and current.status == 0

    test_client.portal.call(submit_stale)


def test_purchase_stock_version_conflict_rolls_back(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
    """库存版本条件更新失败时，不留下新建余额或已确认状态。"""
    document = _create(test_client, auth_headers, purchase_body)
    original = StockPostingService._lock_balance

    async def change_version(self: StockPostingService, warehouse_id: int, goods_id: int, *, create: bool) -> StockBalanceModel:
        """锁后注入版本变化，验证库存 CAS 防护及事务回滚。"""
        row = await original(self, warehouse_id, goods_id, create=create)
        await self.db.execute(update(StockBalanceModel).where(StockBalanceModel.id == row.id).values(version=row.version + 1).execution_options(synchronize_session=False))
        return row

    monkeypatch.setattr(StockPostingService, "_lock_balance", change_version)
    response = test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0})
    assert response.status_code == 409, response.text
    assert _accounting(test_client, purchase_body) == {"balances": [], "logs": [], "bills": [], "audit": []}
    current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
    assert current["status"] == 0 and current["version"] == 0


def test_purchase_returnable_rejects_over_return(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """历史退货超量时明确报错，不将负数静默截断以掩盖账目异常。"""
    document = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    _seed_return(test_client, document, 1, "1.001")
    result = test_client.get(BASE + f"/returnable/{document['id']}", headers=auth_headers)
    assert result.status_code == 409, result.text
