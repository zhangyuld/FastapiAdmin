"""销售报表当前有效业务、跨月退货、快照、导出和独立授权回归。"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import async_db_session
from app.modules.inventory.basic.model import CustomerModel, GoodsCategoryModel, GoodsModel, WarehouseModel
from app.modules.inventory.sale.model import SaleOutItemModel, SaleOutModel, SaleReturnItemModel, SaleReturnModel
from app.modules.system.user.model import UserModel

PREFIX = "/inventory/report/sale"


@pytest.fixture
def sale_data(test_client, auth_headers):
    """只在临时 SQLite 建立明细，包含同商品多行、跨月退货及无效业务。"""
    key = uuid4().hex[:12]

    async def seed():
        async with async_db_session() as db, db.begin():
            user_id = await db.scalar(select(UserModel.id).where(UserModel.username == "super"))
            customer = CustomerModel(customer_code=key, customer_name=key + "Customer", status=1, created_id=user_id)
            categories = [GoodsCategoryModel(name=key + "Category" + str(i)) for i in range(2)]
            warehouses = [WarehouseModel(warehouse_code=key + str(i), warehouse_name=key + "Warehouse" + str(i), status=1) for i in range(2)]
            db.add_all([customer, *categories, *warehouses])
            await db.flush()
            goods = [GoodsModel(goods_code=key + ("G0" if i == 0 else "%_G1"), goods_name=key + "Goods" + str(i), category_id=categories[i].id, unit="件" if i == 0 else "箱") for i in range(2)]
            db.add_all(goods)
            await db.flush()
            originals = []

            async def document(is_return, day, status, warehouse, lines, source=None, deleted=False):
                """构造真实主从结构，返回金额以业务明细快照固化。"""
                model = SaleReturnModel if is_return else SaleOutModel
                item_model = SaleReturnItemModel if is_return else SaleOutItemModel
                values = {
                    "order_no": key + ("PR" if is_return else "PI") + uuid4().hex[:12],
                    "customer_id": customer.id,
                    "warehouse_id": warehouses[warehouse].id,
                    "business_date": date.fromisoformat(day),
                    "status": status,
                    "total_amount": sum(Decimal(amount) for _, _, amount, _ in lines),
                    "created_id": user_id,
                    "is_deleted": deleted,
                }
                if is_return:
                    values.update(source_sale_out_id=source.id, return_reason="report-test")
                total = values.pop("total_amount")
                if is_return:
                    values["return_amount"] = total
                else:
                    values.update(original_amount=total, discount_amount=Decimal("0"), receivable_amount=total)
                row = model(**values)
                db.add(row)
                await db.flush()
                created_lines = []
                for index, (g, quantity, amount, unit) in enumerate(lines, 1):
                    values = {
                        "line_no": index,
                        "goods_id": goods[g].id,
                        "goods_code": goods[g].goods_code,
                        "goods_name": goods[g].goods_name,
                        "spec": "snapshot",
                        "unit": unit,
                        "quantity": Decimal(quantity),
                        "unit_price": Decimal(amount) / Decimal(quantity),
                        "line_amount": Decimal(amount),
                        "created_id": user_id,
                    }
                    if is_return:
                        values.update(sale_return_id=row.id, unit_cost=Decimal(amount) / Decimal(quantity), cost_amount=Decimal(amount), source_sale_out_item_id=originals[0][1][0].id)
                    else:
                        values.update(sale_out_id=row.id, unit_cost=Decimal("1"), cost_amount=Decimal(quantity))
                    item = item_model(**values)
                    db.add(item)
                    created_lines.append(item)
                await db.flush()
                return row, created_lines

            originals.append(await document(False, "2026-09-30", 1, 0, [(0, "10", "100", "件"), (1, "2", "200", "箱"), (0, "2", "20", "件")]))
            originals.append(await document(False, "2026-10-05", 1, 0, [(0, "3", "30", "箱")]))
            originals.append(await document(False, "2026-10-06", 1, 1, [(0, "1", "10", "件")]))
            returned, _ = await document(True, "2026-10-10", 1, 0, [(0, "3", "30", "件")], source=originals[0][0])
            await document(True, "2026-10-11", 1, 1, [(0, "1", "10", "件")], source=originals[2][0])
            for status in [0, 2]:
                await document(False, "2026-10-05", status, 0, [(0, "1", "999", "件")])
                await document(True, "2026-10-12", status, 0, [(0, "1", "10", "件")], source=originals[0][0])
            await document(False, "2026-10-05", 1, 0, [(0, "1", "777", "件")], deleted=True)
            deleted_doc, deleted_lines = await document(False, "2026-10-05", 1, 0, [(0, "1", "888", "件")])
            deleted_lines[0].is_deleted = True
            goods[0].goods_code = key + "CURRENT"
            goods[0].goods_name = "renamed"
            customer.is_deleted = True
            return {
                "key": key,
                "customer_name": customer.customer_name,
                "warehouse_ids": [row.id for row in warehouses],
                "category_ids": [row.id for row in categories],
                "goods_codes": [key + "G0", key + "%_G1"],
                "original_id": originals[0][0].id,
                "return_id": returned.id,
            }

    with patch("app.core.router_class._write_operation_log_async", new=AsyncMock()), patch("app.modules.system.auth.service._write_login_log", new=AsyncMock()):
        yield test_client.portal.call(seed)


def data(client, headers, fixture, endpoint, params=None):
    """所有断言限定本测试客户，避免其他测试已确认单据干扰。"""
    query = {"customer_name": fixture["customer_name"], **(params or {})}
    response = client.get(PREFIX + endpoint, headers=headers, params=query)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_full_summary_and_document_identity(test_client, auth_headers, sale_data):
    """有效明细不重复计整单，多行出库只算一张，两个业务 ID 不碰撞。"""
    summary = data(test_client, auth_headers, sale_data, "/summary")
    assert summary == {"document_count": 5, "sale_amount": "360.00", "return_amount": "40.00", "net_sale_amount": "320.00"}
    documents = data(test_client, auth_headers, sale_data, "/list")
    assert documents["total"] == 5
    assert len({row["id"] for row in documents["items"]}) == 5
    returned = next(row for row in documents["items"] if row["source"] == {"business_type": "sale_return", "business_id": sale_data["return_id"]})
    assert returned["sale_amount"] == "0.00" and returned["return_amount"] == "30.00" and returned["net_sale_amount"] == "-30.00"


def test_cross_month_return_and_negative_period(test_client, auth_headers, sale_data):
    """退货计入其自身月份，即使原销售在期外也保留负向期间金额。"""
    september = data(test_client, auth_headers, sale_data, "/summary", {"business_date": ["2026-09-01", "2026-09-30"]})
    assert september["sale_amount"] == "320.00" and september["return_amount"] == "0.00"
    october = data(test_client, auth_headers, sale_data, "/summary", {"business_date": ["2026-10-01", "2026-10-31"]})
    assert october["document_count"] == 4 and october["net_sale_amount"] == "0.00"
    returned = data(test_client, auth_headers, sale_data, "/summary", {"business_date": ["2026-10-10", "2026-10-10"]})
    assert returned["sale_amount"] == "0.00" and returned["net_sale_amount"] == "-30.00"


def test_goods_filter_uses_lines_and_historical_snapshots(test_client, auth_headers, sale_data):
    """只筛选命中商品，不把另一个商品金额带入，档案改名不改变快照。"""
    summary = data(test_client, auth_headers, sale_data, "/summary", {"goods_code": sale_data["goods_codes"][0]})
    assert summary["sale_amount"] == "160.00" and summary["return_amount"] == "40.00"
    summary = data(test_client, auth_headers, sale_data, "/summary", {"goods_code": "%"})
    assert summary["sale_amount"] == "200.00" and summary["document_count"] == 1
    summary = data(test_client, auth_headers, sale_data, "/summary", {"category_id": sale_data["category_ids"][1]})
    assert summary["sale_amount"] == "200.00"


def test_goods_groups_and_stable_paging(test_client, auth_headers, sale_data):
    """单位不同不合并，同商品同仓多行正确累加，分页 ID 不依赖位置。"""
    goods = data(test_client, auth_headers, sale_data, "/list", {"view": "goods"})
    assert goods["total"] == 4
    row = next(row for row in goods["items"] if row["goods_code"] == sale_data["goods_codes"][0] and row["warehouse_name"].endswith("0") and row["unit"] == "件")
    assert (row["sale_quantity"], row["return_quantity"], row["net_quantity"]) == ("12.000", "3.000", "9.000")
    assert row["net_sale_amount"] == "90.00"
    ids = [data(test_client, auth_headers, sale_data, "/list", {"view": "goods", "page_no": page, "page_size": 1})["items"][0]["id"] for page in range(1, 5)]
    assert set(ids) == {row["id"] for row in goods["items"]}
    summary = data(test_client, auth_headers, sale_data, "/summary", {"view": "goods", "page_size": 1})
    assert summary["net_sale_amount"] == "320.00"


def test_name_and_id_filters_and_empty(test_client, auth_headers, sale_data):
    """名称与 ID 过滤同时支持，空结果返回真实零汇总和空列表。"""
    filtered = data(test_client, auth_headers, sale_data, "/summary", {"warehouse_id": sale_data["warehouse_ids"][1]})
    assert filtered["document_count"] == 2 and filtered["sale_amount"] == "10.00"
    filtered = data(test_client, auth_headers, sale_data, "/summary", {"warehouse_name": sale_data["key"] + "Warehouse1"})
    assert filtered["return_amount"] == "10.00"
    filtered = data(test_client, auth_headers, sale_data, "/summary", {"category_name": sale_data["key"] + "Category1"})
    assert filtered["sale_amount"] == "200.00"
    empty = data(test_client, auth_headers, sale_data, "/summary", {"goods_name": "unmatched"})
    assert empty == {"document_count": 0, "sale_amount": "0.00", "return_amount": "0.00", "net_sale_amount": "0.00"}


@pytest.mark.parametrize(
    "params",
    [
        {"view": "balance"},
        {"warehouse_id": 0},
        {"business_date": ["2026-10-13", "2026-10-01"]},
        {"business_date": ["2026-10-13"]},
        {"view": "goods", "business_no": "PI"},
        {"order_by": '[{"no_such_column":"asc"}]'},
        {"order_by": '[{"id":"drop"}]'},
        {"order_by": "[{}]"},
    ],
)
def test_invalid_query(test_client, auth_headers, sale_data, params):
    """视图、区间和排序在后端校验，不依赖前端白名单。"""
    assert test_client.get(PREFIX + "/list", headers=auth_headers, params={"customer_name": sale_data["customer_name"], **params}).status_code == 422


def test_exports_follow_view_fields_order_and_filename(test_client, auth_headers, sale_data):
    """CSV 和 Excel 使用完整筛选与用户字段顺序，空文件名自动生成。"""
    body = {"query": {"customer_name": sale_data["customer_name"]}, "fields": ["net_sale_amount", "business_no"], "format": "csv", "filename": "核对.xlsx"}
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "净销售金额,来源单号" and len(lines) == 6
    assert "%E6%A0%B8%E5%AF%B9.csv" in response.headers["content-disposition"]
    body.update(query={"view": "goods", "customer_name": sale_data["customer_name"]}, fields=["goods_name", "net_quantity", "net_sale_amount"], format="xlsx", filename="")
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 5 and sheet.max_column == 3
    assert response.headers["content-disposition"].endswith(".xlsx")
    body["fields"] = ["business_no"]
    assert test_client.post(PREFIX + "/export", headers=auth_headers, json=body).status_code == 422
    body.update(fields=["goods_name"], filename="../bad")
    assert test_client.post(PREFIX + "/export", headers=auth_headers, json=body).status_code == 422


def test_readonly_and_source_permissions(test_client, auth_headers, sale_data):
    """报表用户读取企业完整数据，但不继承销售源单的数据范围和导出权限。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.dependencies import get_current_user

    async def load_user():
        async with async_db_session() as db:
            user = await db.scalar(select(UserModel).where(UserModel.username == "user"))
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    auth = AuthSchema(user=test_client.portal.call(load_user), permissions=["module_inventory:sale_report:query", "module_inventory:sale_out:detail"])
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        assert data(test_client, {}, sale_data, "/summary")["net_sale_amount"] == "320.00"
        assert test_client.get("/inventory/sale/out/detail/" + str(sale_data["original_id"])).status_code == 404
        assert test_client.post(PREFIX + "/export", json={"fields": ["business_no"]}).status_code == 403
        auth.permissions = ["module_inventory:sale_out:query"]
        assert test_client.get(PREFIX + "/list").status_code == 403
        assert test_client.get(PREFIX + "/summary").status_code == 403
        for endpoint in ["create", "update/1", "delete"]:
            assert test_client.post(PREFIX + "/" + endpoint, json={}).status_code in (404, 405)
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous


def test_void_recalculates_effective_history(test_client, auth_headers, sale_data):
    """作废移除原业务贡献，不在作废日期再扣一遍，也不保留已作废退货。"""

    async def void_return():
        async with async_db_session() as db, db.begin():
            row = await db.get(SaleReturnModel, sale_data["return_id"])
            row.status = 2

    test_client.portal.call(void_return)
    assert data(test_client, auth_headers, sale_data, "/summary")["net_sale_amount"] == "350.00"

    async def void_original():
        async with async_db_session() as db, db.begin():
            row = await db.get(SaleOutModel, sale_data["original_id"])
            row.status = 2

    test_client.portal.call(void_original)
    assert data(test_client, auth_headers, sale_data, "/summary", {"business_date": ["2026-09-01", "2026-09-30"]})["document_count"] == 0
    assert data(test_client, auth_headers, sale_data, "/summary")["net_sale_amount"] == "30.00"


def test_document_number_with_goods_filter(test_client, auth_headers, sale_data):
    """单号和商品交集只计入该单命中的明细，不影响来源指针。"""
    docs = data(test_client, auth_headers, sale_data, "/list")["items"]
    original = next(row for row in docs if row["source"] == {"business_type": "sale_out", "business_id": sale_data["original_id"]})
    result = data(test_client, auth_headers, sale_data, "/list", {"business_no": original["business_no"], "goods_code": sale_data["goods_codes"][0]})
    assert result["total"] == 1 and result["items"][0]["sale_amount"] == "120.00"


def test_export_text_is_not_formula(test_client, auth_headers, sale_data):
    """档案公式文本导出为字面值，CSV 和 Excel 都不执行，负金额不加引号。"""
    import csv
    from io import StringIO

    from sqlalchemy import update

    async def rename_customer():
        async with async_db_session() as db, db.begin():
            await db.execute(update(CustomerModel).where(CustomerModel.customer_code == sale_data["key"]).values(customer_name="=2+2"))

    test_client.portal.call(rename_customer)
    body = {"query": {"warehouse_name": sale_data["key"]}, "fields": ["customer_name", "net_sale_amount"], "format": "csv"}
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200
    values = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
    assert all(row[0] == "'=2+2" for row in values[1:])
    assert any(row[1] == "-30.00" for row in values[1:])
    body["format"] = "xlsx"
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    cell = load_workbook(BytesIO(response.content)).active.cell(2, 1)
    assert cell.value == "'=2+2" and cell.data_type == "s"


def test_export_over_limit_is_explicit(test_client):
    """框架最多导出10万行，超过上限必须报错而不是静默截断。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.exceptions import CustomException
    from app.modules.inventory.report.sale.schema import SaleReportExport
    from app.modules.inventory.report.sale.service import SaleReportService

    async def run():
        async with async_db_session() as db:
            auth = AuthSchema(user=CoreUserSchema(id=1, username="super", name="super", dept_id=None, is_superuser=True))
            service = SaleReportService(auth, db)
            with patch.object(service.crud, "rows", new=AsyncMock(return_value=([], 100001))):
                with pytest.raises(CustomException, match="10万行"):
                    await service.export(SaleReportExport(fields=["business_no"]))

    test_client.portal.call(run)


def test_discounted_sale_and_stored_return_amount(test_client, auth_headers, sale_data):
    """优惠后销售与固化退货分别汇总，不按原价重复估算退货。"""

    async def change_amounts():
        async with async_db_session() as db, db.begin():
            original = await db.get(SaleOutModel, sale_data["original_id"])
            original.discount_amount = Decimal("32.00")
            original.receivable_amount = Decimal("288.00")
            returned = await db.get(SaleReturnModel, sale_data["return_id"])
            returned.return_amount = Decimal("27.00")
            item = await db.scalar(select(SaleReturnItemModel).where(SaleReturnItemModel.sale_return_id == returned.id))
            item.line_amount = Decimal("27.00")

    test_client.portal.call(change_amounts)
    summary = data(test_client, auth_headers, sale_data, "/summary")
    assert summary == {"document_count": 5, "sale_amount": "328.00", "return_amount": "37.00", "net_sale_amount": "291.00"}
    filtered = data(test_client, auth_headers, sale_data, "/summary", {"goods_code": sale_data["goods_codes"][0]})
    assert filtered["sale_amount"] == "148.00" and filtered["net_sale_amount"] == "111.00"
    response = test_client.post(
        PREFIX + "/export",
        headers=auth_headers,
        json={
            "query": {"customer_name": sale_data["customer_name"]},
            "fields": ["customer_name", "sale_amount", "return_amount", "net_sale_amount"],
            "format": "csv",
        },
    )
    assert response.status_code == 200, response.text
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "客户,销售金额,退货金额,净销售金额"
    assert any(line.endswith(",288.00,0.00,288.00") for line in lines[1:])
    assert any(line.endswith(",0.00,27.00,-27.00") for line in lines[1:])


def test_mysql_grouping_keeps_case_sensitive_snapshot_keys():
    """MySQL 编译使用二进制快照键，不被默认大小写/重音排序规则合并。"""
    from sqlalchemy.dialects import mysql

    from app.modules.inventory.report.sale.crud import SaleReportCRUD
    from app.modules.inventory.report.sale.schema import SaleReportQuery

    lines = SaleReportCRUD._branch(SaleReportQuery(view="goods"), False).subquery()
    grouped = SaleReportCRUD._grouped(lines, "goods")
    sql = str(select(grouped).compile(dialect=mysql.dialect()))
    assert "AS BINARY" in sql
    assert "snapshot_unit" in sql and "snapshot_goods_name" in sql
    assert "CAST" in sql.split("GROUP BY", 1)[1]


@pytest.mark.parametrize(
    "amounts,discount",
    [
        (["100.00", "200.00", "20.00"], "32.00"),
        (["0.01", "0.01", "0.01"], "0.01"),
        (["0.01", "0.01", "0.00"], "0.01"),
        (["0.00", "0.00", "0.00"], "0.00"),
        (["100.00", "200.00", "20.00"], "320.00"),
        (["0.00", "1.03", "2.07"], "0.43"),
        (["17.99", "0.07", "34.56"], "13.37"),
    ],
)
def test_discount_allocation_matches_return_rules(test_client, auth_headers, sale_data, amounts, discount):
    """逐行比对原退货分摊规则，商品过滤不能缩小整单分摊分母。"""
    from app.modules.inventory.report.sale.crud import SaleReportCRUD
    from app.modules.inventory.sale.out.amount import allocate_receivable

    expected = allocate_receivable([Decimal(value) for value in amounts], Decimal(discount))

    async def change_amounts():
        async with async_db_session() as db, db.begin():
            document = await db.get(SaleOutModel, sale_data["original_id"])
            items = list((await db.scalars(select(SaleOutItemModel).where(SaleOutItemModel.sale_out_id == document.id).order_by(SaleOutItemModel.line_no))).all())
            for item, amount in zip(items, amounts, strict=True):
                item.line_amount = Decimal(amount)
            document.original_amount = sum(map(Decimal, amounts))
            document.discount_amount = Decimal(discount)
            document.receivable_amount = document.original_amount - document.discount_amount
            await db.flush()
            allocated = SaleReportCRUD._allocated()
            values = list((await db.scalars(select(allocated.c.sale_amount).where(allocated.c.id.in_([item.id for item in items])).order_by(allocated.c.id))).all())
            assert values == expected

    test_client.portal.call(change_amounts)
    period = {"business_date": ["2026-09-01", "2026-09-30"]}
    summary = data(test_client, auth_headers, sale_data, "/summary", period)
    assert summary["sale_amount"] == format(sum(expected), ".2f")
    filtered = data(test_client, auth_headers, sale_data, "/summary", {**period, "goods_code": sale_data["goods_codes"][0]})
    assert filtered["sale_amount"] == format(expected[0] + expected[2], ".2f")
    documents = data(test_client, auth_headers, sale_data, "/list", {**period, "goods_code": sale_data["goods_codes"][1]})
    assert documents["items"][0]["sale_amount"] == format(expected[1], ".2f")


def test_sql_allocation_matches_random_cent_amounts(test_client, auth_headers, sale_data):
    """固定随机样本比对累计 HALF_UP，覆盖连续行的分摊尾差。"""
    import random

    from app.modules.inventory.report.sale.crud import SaleReportCRUD
    from app.modules.inventory.sale.out.amount import allocate_receivable

    async def run():
        randomizer = random.Random(20261013)
        async with async_db_session() as db, db.begin():
            document = await db.get(SaleOutModel, sale_data["original_id"])
            items = list((await db.scalars(select(SaleOutItemModel).where(SaleOutItemModel.sale_out_id == document.id).order_by(SaleOutItemModel.line_no))).all())
            for _ in range(60):
                amounts = [Decimal(randomizer.randrange(10000)) / 100 for _ in items]
                discount = Decimal(randomizer.randrange(int(sum(amounts) * 100) + 1)) / 100
                for item, amount in zip(items, amounts, strict=True):
                    item.line_amount = amount
                document.original_amount = sum(amounts)
                document.discount_amount = discount
                document.receivable_amount = sum(amounts) - discount
                await db.flush()
                allocated = SaleReportCRUD._allocated()
                actual = list((await db.scalars(select(allocated.c.sale_amount).where(allocated.c.id.in_([item.id for item in items])).order_by(allocated.c.id))).all())
                assert actual == allocate_receivable(amounts, discount), (amounts, discount, actual)

    test_client.portal.call(run)
