import csv
import re
from copy import deepcopy
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, select, update
from test_inventory_purchase_in import _accounting, _confirm, _create
from test_inventory_purchase_in import disable_purchase_test_audit_writes as disable_purchase_test_audit_writes
from test_inventory_purchase_in import purchase_body as _purchase_body_fixture

from app.core.base_schema import AuthSchema, CoreUserSchema
from app.core.database import async_db_session
from app.core.dependencies import get_current_user
from app.core.exceptions import CustomException
from app.modules.inventory.basic.model import GoodsModel
from app.modules.inventory.common.business_log import BusinessLogService
from app.modules.inventory.common.model import BusinessLogModel
from app.modules.inventory.finance.model import PayableModel
from app.modules.inventory.finance.payable import PayableService
from app.modules.inventory.purchase.model import PurchaseInModel, PurchaseReturnItemModel, PurchaseReturnModel
from app.modules.inventory.purchase.return_order import crud as crud_module
from app.modules.inventory.stock.model import StockBalanceModel, StockLogModel
from app.modules.inventory.stock.posting import StockPostingService
from app.modules.system.user.model import UserModel

# 复用入库测试的档案准备，不修改全局 conftest 或开发数据库。
purchase_body = _purchase_body_fixture

BASE = "/inventory/purchase/return_order"
IN_BASE = "/inventory/purchase/inbound"


@pytest.fixture
def source_document(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> dict[str, Any]:
    """创建已确认来源，两种商品初始数量十件、三件，金额二十元、十二元。"""
    body = deepcopy(purchase_body)
    body["items"][0].update(quantity="10.000", unit_price="2.0000")
    body["items"][1].update(quantity="3.000", unit_price="4.0000")
    return _confirm(test_client, auth_headers, _create(test_client, auth_headers, body))


@pytest.fixture
def return_body(source_document: dict[str, Any]) -> dict[str, Any]:
    """构建两行原单退货请求，总金额八元。"""
    return {
        "supplier_id": source_document["supplier_id"],
        "warehouse_id": source_document["warehouse_id"],
        "source_purchase_in_id": source_document["id"],
        "business_date": "2026-10-12",
        "return_reason": "  质量问题  ",
        "remark": "退货草稿",
        "items": [
            {
                "goods_id": item["goods_id"],
                "source_purchase_in_item_id": item["id"],
                "quantity": "2.000" if index == 0 else "1.000",
                "unit_price": item["unit_price"],
            }
            for index, item in enumerate(source_document["items"])
        ],
    }


def _save(client: TestClient, headers: dict[str, str], body: dict[str, Any]) -> dict[str, Any]:
    """保存退货草稿并读取完整详情。"""
    result = client.post(BASE + "/create", headers=headers, json=body)
    assert result.status_code == 201, result.text
    return result.json()["data"]


def _post_confirm(client: TestClient, headers: dict[str, str], document: dict[str, Any]) -> dict[str, Any]:
    """确认当前版本的退货。"""
    result = client.post(BASE + f"/confirm/{document['id']}", headers=headers, json={"version": document["version"]})
    assert result.status_code == 200, result.text
    return result.json()["data"]


def _void(client: TestClient, headers: dict[str, str], document: dict[str, Any]) -> dict[str, Any]:
    """按版本作废已确认退货。"""
    result = client.post(BASE + f"/void/{document['id']}", headers=headers, json={"version": document["version"], "void_reason": "退货取消"})
    assert result.status_code == 200, result.text
    return result.json()["data"]


def _snapshot(client: TestClient, body: dict[str, Any]) -> dict[str, Any]:
    """记账快照包含来源与退货全部审计，用于检测事务残留。"""
    result = _accounting(client, body)

    async def audit() -> list[tuple[Any, ...]]:
        """读取本测试来源下的退货业务日志。"""
        async with async_db_session() as db:
            ids = select(PurchaseReturnModel.id).where(PurchaseReturnModel.source_purchase_in_id == body["source_purchase_in_id"])
            return (
                await db.execute(
                    select(BusinessLogModel.action, BusinessLogModel.before_status, BusinessLogModel.after_status)
                    .where(
                        BusinessLogModel.business_type == "purchase_return",
                        BusinessLogModel.business_id.in_(ids),
                    )
                    .order_by(BusinessLogModel.id)
                )
            ).all()

    result["return_audit"] = client.portal.call(audit)
    return result


def _available(client: TestClient, headers: dict[str, str], source_id: int) -> list[str]:
    """查询来源行实时可退数量。"""
    result = client.get(IN_BASE + f"/returnable/{source_id}", headers=headers)
    assert result.status_code == 200, result.text
    return [item["returnable_quantity"] for item in result.json()["data"]["items"]]


def test_source_detail_return_summary_and_history(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any]) -> None:
    """原单展示未退、部分、全部和作废恢复，历史草稿不影响汇总。"""
    source_id = source_document["id"]

    def detail() -> dict[str, Any]:
        """重新读取原入库单实际退货状态。"""
        response = test_client.get(IN_BASE + f"/detail/{source_id}", headers=auth_headers)
        assert response.status_code == 200, response.text
        return response.json()["data"]

    draft = _save(test_client, auth_headers, return_body)
    assert detail()["return_summary"] == {"return_status": "none", "confirmed_count": 0, "returned_amount": "0.00"}
    first = _post_confirm(test_client, auth_headers, draft)
    current = detail()
    assert current["return_summary"] == {"return_status": "partial", "confirmed_count": 1, "returned_amount": "8.00"}
    assert [line["returned_quantity"] for line in current["items"]] == ["2.000", "1.000"]
    assert [line["returnable_quantity"] for line in current["items"]] == ["8.000", "2.000"]
    assert current["total_amount"] == source_document["total_amount"]
    remainder = deepcopy(return_body)
    for line, quantity in zip(remainder["items"], ["8.000", "2.000"], strict=True):
        line["quantity"] = quantity
    second = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, remainder))
    assert detail()["return_summary"] == {"return_status": "full", "confirmed_count": 2, "returned_amount": "32.00"}
    _void(test_client, auth_headers, first)
    assert detail()["return_summary"]["return_status"] == "partial"
    _void(test_client, auth_headers, second)
    assert detail()["return_summary"] == {"return_status": "none", "confirmed_count": 0, "returned_amount": "0.00"}
    _save(test_client, auth_headers, return_body)
    listed = test_client.get(IN_BASE + f"/returns/{source_id}", headers=auth_headers, params={"page_no": 1, "page_size": 1})
    assert listed.status_code == 200, listed.text
    page = listed.json()["data"]
    assert page["total"] == 3 and page["has_next"] and len(page["items"]) == 1
    assert page["items"][0]["status"] == 0
    assert test_client.get(IN_BASE + "/returns/999999999", headers=auth_headers).status_code == 404


def test_source_return_history_permissions_and_scope(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], purchase_body: dict[str, Any]) -> None:
    """跨范围退货计入原单汇总，但记录列表和单据详情不越权。"""
    source = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    body = deepcopy(return_body)
    body["source_purchase_in_id"] = source["id"]
    body["items"] = [{
        "goods_id": source["items"][0]["goods_id"], "source_purchase_in_item_id": source["items"][0]["id"],
        "quantity": "0.500", "unit_price": source["items"][0]["unit_price"],
    }]
    returned = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))

    async def prepare_user() -> CoreUserSchema:
        """将原单归属普通用户，保持下游退货归属其他用户。"""
        async with async_db_session() as db, db.begin():
            user = (await db.execute(select(UserModel).where(UserModel.username == "user"))).scalar_one()
            await db.execute(update(PurchaseInModel).where(PurchaseInModel.id == source["id"]).values(created_id=user.id))
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    auth = AuthSchema(user=test_client.portal.call(prepare_user), permissions=["module_inventory:purchase_in:detail"])
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        detail = test_client.get(IN_BASE + f"/detail/{source['id']}").json()["data"]
        assert detail["return_summary"]["return_status"] == "partial"
        assert detail["items"][0]["returned_quantity"] == "0.500"
        assert test_client.get(IN_BASE + f"/returns/{source['id']}").status_code == 403
        auth.permissions.append("module_inventory:purchase_return:query")
        listed = test_client.get(IN_BASE + f"/returns/{source['id']}")
        assert listed.status_code == 200, listed.text
        assert listed.json()["data"]["total"] == 0
        assert test_client.get(IN_BASE + f"/returns/{return_body['source_purchase_in_id']}").status_code == 404
        auth.permissions.append("module_inventory:purchase_return:detail")
        assert test_client.get(BASE + f"/detail/{returned['id']}").status_code == 404
        auth.permissions.remove("module_inventory:purchase_in:detail")
        assert test_client.get(IN_BASE + f"/returns/{source['id']}").status_code == 403
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous


def test_return_draft_lifecycle_and_filter(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any]) -> None:
    """草稿创建、来源筛选、更新及删除不影响库存和应付。"""
    before = _snapshot(test_client, return_body)
    document = _save(test_client, auth_headers, return_body)
    assert document["total_amount"] == "8.00" and document["return_reason"] == "质量问题"
    assert document["source_order_no"] == source_document["order_no"]
    assert document["items"][0]["cost_amount"] == "4.00" and document["items"][0]["unit_cost"] == "2.0000"
    assert _snapshot(test_client, return_body) == before
    assert _available(test_client, auth_headers, source_document["id"]) == ["10.000", "3.000"]
    listed = test_client.get(
        BASE + "/list",
        headers=auth_headers,
        params=[
            ("source_order_no", source_document["order_no"]),
            ("status", "0"),
            ("business_date", "2026-10-12"),
            ("business_date", "2026-10-12"),
            ("created_id", document["created_id"]),
            ("supplier_id", document["supplier_id"]),
        ],
    )
    assert listed.status_code == 200, listed.text
    assert listed.json()["data"]["total"] == 1
    assert "items" not in listed.json()["data"]["items"][0]
    assert test_client.get(BASE + "/list", headers=auth_headers, params={"source_order_no": "nonexistent"}).json()["data"]["total"] == 0
    edited = deepcopy(return_body)
    edited["items"] = edited["items"][:1]
    edited["items"][0]["quantity"] = "1.000"
    edited["version"] = 0
    update_result = test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json=edited)
    assert update_result.status_code == 200, update_result.text
    current = update_result.json()["data"]
    assert current["version"] == 1 and current["total_amount"] == "2.00" and current["order_no"] == document["order_no"]
    assert test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json=edited).status_code == 409
    assert test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [document["id"]], "versions": {document["id"]: 0}}).status_code == 409
    deleted = test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [document["id"]], "versions": {document["id"]: 1}})
    assert deleted.status_code == 200, deleted.text
    assert test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).status_code == 404

    async def deleted_items() -> bool:
        """主从表必须同时软删除。"""
        async with async_db_session() as db:
            items = (await db.execute(select(PurchaseReturnItemModel).where(PurchaseReturnItemModel.purchase_return_id == document["id"]))).scalars().all()
            return all(row.is_deleted for row in items)

    assert test_client.portal.call(deleted_items)
    assert _snapshot(test_client, return_body) == before


def test_return_confirmation_and_void(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any]) -> None:
    """退货扣库存和应付，作废恢复并保留关联反向流水及原记录。"""
    before = _snapshot(test_client, return_body)
    document = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, return_body))
    assert document["status"] == 1 and document["version"] == 1 and document["confirmed_time"] and document["confirmed_name"]
    posted = _snapshot(test_client, return_body)
    assert posted["balances"] == [(Decimal("8"), Decimal("16"), Decimal("2"), 2), (Decimal("2"), Decimal("8"), Decimal("4"), 2)]
    assert posted["bills"][-1][2] == Decimal("-8")
    assert posted["return_audit"] == [("confirm", 0, 1)]
    assert _available(test_client, auth_headers, source_document["id"]) == ["8.000", "2.000"]
    assert test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0}).status_code == 409
    assert test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json={**return_body, "version": 1}).status_code == 409
    assert test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [document["id"]], "versions": {document["id"]: 1}}).status_code == 409
    assert test_client.post(IN_BASE + f"/void/{source_document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "应被阻止"}).status_code == 409
    assert _snapshot(test_client, return_body) == posted
    voided = _void(test_client, auth_headers, document)
    assert voided["status"] == 2 and voided["version"] == 2 and voided["voided_time"] and voided["voided_name"]
    restored = _snapshot(test_client, return_body)
    assert [row[:3] for row in restored["balances"]] == [row[:3] for row in before["balances"]]
    assert restored["logs"][:4] == posted["logs"] and restored["bills"][:2] == posted["bills"]
    originals = {row[1]: row for row in posted["logs"][2:]}
    for row in restored["logs"][4:]:
        assert row[3:5] == (-originals[row[1]][3], -originals[row[1]][4])
        assert row[5] == originals[row[1]][0] and not row[6]
    assert restored["bills"][-1][2] == Decimal("8") and restored["bills"][-1][5] == posted["bills"][-1][0]
    assert restored["return_audit"] == [("confirm", 0, 1), ("void", 1, 2)]
    assert _available(test_client, auth_headers, source_document["id"]) == ["10.000", "3.000"]
    assert test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "重复"}).status_code == 409
    source_void = test_client.post(IN_BASE + f"/void/{source_document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "退货已作废"})
    assert source_void.status_code == 200, source_void.text


def test_return_multiple_drafts_do_not_reserve_and_confirm_rechecks(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any]) -> None:
    """多个草稿可同时录入相同额度，但第二次确认不能超出实时剩余数量。"""
    return_body["items"] = return_body["items"][:1]
    return_body["items"][0]["quantity"] = "6"
    first = _save(test_client, auth_headers, return_body)
    second = _save(test_client, auth_headers, return_body)
    _post_confirm(test_client, auth_headers, first)
    before = _snapshot(test_client, return_body)
    result = test_client.post(BASE + f"/confirm/{second['id']}", headers=auth_headers, json={"version": 0})
    assert result.status_code == 409, result.text
    assert _snapshot(test_client, return_body) == before
    assert _available(test_client, auth_headers, source_document["id"]) == ["4.000", "3.000"]


@pytest.mark.parametrize(
    "field,value",
    [
        ("source_purchase_in_id", None),
        ("return_reason", ""),
        ("total_amount", "0"),
    ],
)
def test_return_rejects_invalid_http_schema(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], field: str, value: Any) -> None:
    """无原单退货和客户端金额注入在 HTTP 入口即被拒绝。"""
    return_body[field] = value
    result = test_client.post(BASE + "/create", headers=auth_headers, json=return_body)
    assert result.status_code == 422, result.text


@pytest.mark.parametrize("tamper", ["supplier", "warehouse", "source", "goods", "price", "quantity"])
def test_return_rejects_tampered_source(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any], tamper: str) -> None:
    """来源关系、商品、价格和可退上限都由后端核对。"""
    if tamper in {"supplier", "warehouse"}:
        return_body[tamper + "_id"] += 99999
    elif tamper == "source":
        return_body["items"] = return_body["items"][:1]
        return_body["items"][0]["source_purchase_in_item_id"] = source_document["items"][1]["id"]
    elif tamper == "goods":
        return_body["items"][0]["goods_id"] = source_document["items"][1]["goods_id"]
    elif tamper == "price":
        return_body["items"][0]["unit_price"] = "0"
    else:
        return_body["items"][0]["quantity"] = "10.001"
    before = _snapshot(test_client, return_body)
    response = test_client.post(BASE + "/create", headers=auth_headers, json=return_body)
    assert response.status_code == (409 if tamper == "quantity" else 400), response.text
    assert _snapshot(test_client, return_body) == before


@pytest.mark.parametrize("state", [0, 2])
def test_return_rejects_non_confirmed_source(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any], state: int) -> None:
    """来源草稿或作废单不能作为退货来源。"""

    async def set_state() -> None:
        """只预置来源状态以覆盖错误分支。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(PurchaseInModel).where(PurchaseInModel.id == source_document["id"]).values(status=state))

    test_client.portal.call(set_state)
    result = test_client.post(BASE + "/create", headers=auth_headers, json=return_body)
    assert result.status_code == 409, result.text


@pytest.mark.parametrize("quantity,amount", [("1", "20"), ("10", "1"), ("2", "5")])
def test_return_confirmation_stock_guards(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], quantity: str, amount: str) -> None:
    """库存数量、金额不足或清空数量后留金额时，状态和流水不生效。"""
    document = _save(test_client, auth_headers, return_body)

    async def change_stock() -> None:
        """模拟来源入库后发生其他库存业务。"""
        async with async_db_session() as db, db.begin():
            await db.execute(
                update(StockBalanceModel)
                .where(StockBalanceModel.warehouse_id == return_body["warehouse_id"], StockBalanceModel.goods_id == return_body["items"][0]["goods_id"])
                .values(quantity=Decimal(quantity), stock_amount=Decimal(amount))
            )

    test_client.portal.call(change_stock)
    before = _snapshot(test_client, return_body)
    result = test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0})
    assert result.status_code == 409, result.text
    assert _snapshot(test_client, return_body) == before
    current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
    assert current["status"] == 0 and current["version"] == 0


@pytest.mark.parametrize("operation", ["confirm", "void"])
@pytest.mark.parametrize("stage", ["stock", "payable", "audit"])
def test_return_posting_rollback(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], monkeypatch: pytest.MonkeyPatch, operation: str, stage: str) -> None:
    """库存、应付或日志写入后失败，主从表金额、状态和全部账目必须回滚。"""
    document = _save(test_client, auth_headers, return_body)
    if operation == "void":
        document = _post_confirm(test_client, auth_headers, document)
    before = _snapshot(test_client, return_body)
    cls, method = {
        "stock": (StockPostingService, "post"),
        "payable": (PayableService, "reverse" if operation == "void" else "create"),
        "audit": (BusinessLogService, "record"),
    }[stage]
    original = getattr(cls, method)

    async def fail(self: Any, *args: Any, **kwargs: Any) -> None:
        """在真实写入后注入业务异常。"""
        await original(self, *args, **kwargs)
        raise CustomException(msg="模拟退货记账失败", status_code=409)

    monkeypatch.setattr(cls, method, fail)
    payload = {"version": document["version"]}
    if operation == "void":
        payload["void_reason"] = "测试冲销"
    result = test_client.post(BASE + f"/{operation}/{document['id']}", headers=auth_headers, json=payload)
    assert result.status_code == 409, result.text
    assert _snapshot(test_client, return_body) == before
    current = test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"]
    assert current == document


def test_return_export_fields_and_filename(test_client: TestClient, auth_headers: dict[str, str], source_document: dict[str, Any], return_body: dict[str, Any]) -> None:
    """服务端导出处理来源筛选、选定字段顺序及自定义和默认文件名。"""
    document = _save(test_client, auth_headers, return_body)
    response = test_client.post(
        BASE + "/export",
        headers=auth_headers,
        json={
            "query": {"source_order_no": source_document["order_no"]},
            "fields": ["source_order_no", "total_amount", "order_no", "status"],
            "format": "xlsx",
            "filename": "退货记录",
        },
    )
    assert response.status_code == 200, response.text
    assert "退货记录.xlsx" in unquote(response.headers["content-disposition"])
    book = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(book["采购退货"].iter_rows(values_only=True))
    book.close()
    assert rows == [("原入库单号", "总金额", "单据编号", "状态"), (source_document["order_no"], "8.00", document["order_no"], "草稿")]
    csv_result = test_client.post(
        BASE + "/export",
        headers=auth_headers,
        json={
            "query": {"order_no": document["order_no"]},
            "fields": ["order_no", "return_reason"],
            "format": "csv",
            "filename": "",
        },
    )
    assert csv_result.status_code == 200, csv_result.text
    assert re.search(r"采购退货列表_\d{8}_\d{6}\.csv", unquote(csv_result.headers["content-disposition"]))
    assert list(csv.reader(StringIO(csv_result.content.decode("utf-8-sig")))) == [["单据编号", "退货原因"], [document["order_no"], "质量问题"]]


def test_return_tail_rounding_and_reconfirm_after_void(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """三次小额退货全部退完不留尾差，作废首笔后再次退货仍守恒。"""
    purchase_body["items"] = purchase_body["items"][:1]
    purchase_body["items"][0].update(quantity="1.000", unit_price="0.0050")
    source = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    item = source["items"][0]
    body = {
        "supplier_id": source["supplier_id"],
        "warehouse_id": source["warehouse_id"],
        "source_purchase_in_id": source["id"],
        "business_date": "2026-10-12",
        "return_reason": "小额分批退货",
        "items": [{"goods_id": item["goods_id"], "source_purchase_in_item_id": item["id"], "quantity": "0.500", "unit_price": item["unit_price"]}],
    }
    first = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))
    body["items"][0]["quantity"] = "0.250"
    second = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))
    third = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))
    assert [row["total_amount"] for row in [first, second, third]] == ["0.01", "0.00", "0.00"]
    assert _snapshot(test_client, body)["balances"][0][:3] == (Decimal("0"), Decimal("0"), Decimal("0"))
    _void(test_client, auth_headers, first)
    assert _available(test_client, auth_headers, source["id"]) == ["0.500"]
    body["items"][0]["quantity"] = "0.500"
    replacement = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))
    assert replacement["total_amount"] == "0.01" and replacement["items"][0]["cost_amount"] == "0.01"
    assert _snapshot(test_client, body)["balances"][0][:3] == (Decimal("0"), Decimal("0"), Decimal("0"))


def test_return_tail_allocation_after_out_of_order_void(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """作废改变累计分摊后，小量退货不能产生负金额，最终仍退完原金额。"""
    purchase_body["items"] = purchase_body["items"][:1]
    purchase_body["items"][0].update(quantity="3.000", unit_price="0.0033")
    source = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    item = source["items"][0]
    body = {
        "supplier_id": source["supplier_id"],
        "warehouse_id": source["warehouse_id"],
        "source_purchase_in_id": source["id"],
        "business_date": "2026-10-12",
        "return_reason": "乱序作废",
        "items": [{"goods_id": item["goods_id"], "source_purchase_in_item_id": item["id"], "quantity": "1.000", "unit_price": item["unit_price"]}],
    }
    documents = [_post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body)) for _ in range(3)]
    assert [row["total_amount"] for row in documents] == ["0.00", "0.01", "0.00"]
    _void(test_client, auth_headers, documents[0])
    _void(test_client, auth_headers, documents[2])
    body["items"][0]["quantity"] = "0.100"
    small = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))
    assert small["total_amount"] == "0.00"
    body["items"][0]["quantity"] = "1.900"
    last = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, body))
    assert last["total_amount"] == "0.00"
    assert _snapshot(test_client, body)["balances"][0][:3] == (Decimal("0"), Decimal("0"), Decimal("0"))


def test_return_draft_confirmation_reallocates_tail(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any]) -> None:
    """两个草稿金额相同，先确认占用尾差后，后确认单据必须重算并固化实际金额。"""
    purchase_body["items"] = purchase_body["items"][:1]
    purchase_body["items"][0].update(quantity="1.000", unit_price="0.0050")
    source = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    item = source["items"][0]
    body = {
        "supplier_id": source["supplier_id"],
        "warehouse_id": source["warehouse_id"],
        "source_purchase_in_id": source["id"],
        "business_date": "2026-10-12",
        "return_reason": "尾差重算",
        "items": [{"goods_id": item["goods_id"], "source_purchase_in_item_id": item["id"], "quantity": "0.500", "unit_price": item["unit_price"]}],
    }
    first, second = [_save(test_client, auth_headers, body) for _ in range(2)]
    assert first["total_amount"] == second["total_amount"] == "0.01"
    assert _post_confirm(test_client, auth_headers, first)["total_amount"] == "0.01"
    confirmed = _post_confirm(test_client, auth_headers, second)
    assert confirmed["total_amount"] == "0.00" and confirmed["items"][0]["cost_amount"] == "0.00"
    assert confirmed["items"][0]["id"] == second["items"][0]["id"]


def test_return_switch_source_and_snapshot(test_client: TestClient, auth_headers: dict[str, str], purchase_body: dict[str, Any], return_body: dict[str, Any]) -> None:
    """草稿可以更换来源，商品档案变更不改变来源或退货历史快照。"""
    document = _save(test_client, auth_headers, return_body)
    replacement = _confirm(test_client, auth_headers, _create(test_client, auth_headers, purchase_body))
    item = replacement["items"][0]
    update_body = {
        **return_body,
        "source_purchase_in_id": replacement["id"],
        "version": 0,
        "items": [
            {
                "goods_id": item["goods_id"],
                "source_purchase_in_item_id": item["id"],
                "quantity": "0.500",
                "unit_price": item["unit_price"],
            }
        ],
    }
    response = test_client.post(BASE + f"/update/{document['id']}", headers=auth_headers, json=update_body)
    assert response.status_code == 200, response.text
    changed = response.json()["data"]
    assert changed["source_order_no"] == replacement["order_no"] and changed["version"] == 1

    async def rename_goods() -> None:
        """确认不能从当前档案重新采集历史名称和参考成本。"""
        async with async_db_session() as db, db.begin():
            await db.execute(update(GoodsModel).where(GoodsModel.id == item["goods_id"]).values(goods_name="更名后商品", reference_cost_price=Decimal("100")))

    test_client.portal.call(rename_goods)
    confirmed = _post_confirm(test_client, auth_headers, changed)
    assert confirmed["items"][0]["goods_name"] == item["goods_name"] and confirmed["items"][0]["unit_cost"] == item["unit_price"]


@pytest.mark.parametrize("kind", ["supplier", "warehouse", "goods"])
def test_return_confirmation_rechecks_archives(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], kind: str) -> None:
    """草稿保存后档案停用，确认也必须拒绝记账。"""
    document = _save(test_client, auth_headers, return_body)
    if kind == "goods":

        async def disable_goods() -> None:
            """商品当前没有批量状态接口，直接预置测试状态。"""
            async with async_db_session() as db, db.begin():
                await db.execute(update(GoodsModel).where(GoodsModel.id == return_body["items"][0]["goods_id"]).values(status=1))

        test_client.portal.call(disable_goods)
    else:
        result = test_client.post(f"/inventory/basic/{kind}/status/batch", headers=auth_headers, json={"ids": [return_body[kind + "_id"]], "status": 1})
        assert result.status_code == 200, result.text
    before = _snapshot(test_client, return_body)
    response = test_client.post(BASE + f"/confirm/{document['id']}", headers=auth_headers, json={"version": 0})
    assert response.status_code == 400, response.text
    assert _snapshot(test_client, return_body) == before


@pytest.mark.parametrize("missing", ["log", "payable"])
def test_return_void_rejects_missing_posting(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], missing: str) -> None:
    """缺少原库存或应付不能直接用单据金额补造冲销。"""
    document = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, return_body))

    async def remove_original() -> None:
        """预置不完整历史流水。"""
        model = StockLogModel if missing == "log" else PayableModel
        async with async_db_session() as db, db.begin():
            await db.execute(delete(model).where(model.business_type == "purchase_return", model.business_id == document["id"]))

    test_client.portal.call(remove_original)
    before = _snapshot(test_client, return_body)
    response = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "缺失流水"})
    assert response.status_code == 409, response.text
    assert _snapshot(test_client, return_body) == before


def test_return_void_rejects_settled_bill(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any]) -> None:
    """负向退货应付已核销时也必须先撤销核销，不能直接作废。"""
    document = _post_confirm(test_client, auth_headers, _save(test_client, auth_headers, return_body))

    async def settle_bill() -> None:
        """模拟退款核销占用负向应付额度。"""
        async with async_db_session() as db, db.begin():
            await db.execute(
                update(PayableModel)
                .where(PayableModel.business_type == "purchase_return", PayableModel.business_id == document["id"])
                .values(allocated_amount=Decimal("1"), unallocated_amount=Decimal("7"), status=1)
            )

    test_client.portal.call(settle_bill)
    before = _snapshot(test_client, return_body)
    response = test_client.post(BASE + f"/void/{document['id']}", headers=auth_headers, json={"version": 1, "void_reason": "核销占用"})
    assert response.status_code == 409 and "核销" in response.json()["msg"]
    assert _snapshot(test_client, return_body) == before


def test_return_batch_delete_rollback(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
    """批量第二笔删除版本冲突，第一笔软删除必须回滚。"""
    documents = [_save(test_client, auth_headers, return_body) for _ in range(2)]
    original = crud_module.PurchaseReturnCRUD.delete_draft
    calls = 0

    async def fail_second(self: Any, id: int, version: int) -> None:
        """模拟第二个版本条件更新失败。"""
        nonlocal calls
        calls += 1
        if calls == 2:
            raise CustomException(msg="模拟删除版本冲突", status_code=409)
        await original(self, id, version)

    monkeypatch.setattr(crud_module.PurchaseReturnCRUD, "delete_draft", fail_second)
    response = test_client.post(BASE + "/delete", headers=auth_headers, json={"ids": [row["id"] for row in documents], "versions": {row["id"]: 0 for row in documents}})
    assert response.status_code == 409, response.text
    for document in documents:
        assert test_client.get(BASE + f"/detail/{document['id']}", headers=auth_headers).json()["data"] == document


def test_return_permissions_and_source_data_scope(test_client: TestClient, auth_headers: dict[str, str], return_body: dict[str, Any], purchase_body: dict[str, Any]) -> None:
    """有退货按钮权限不代表有他人的退货或来源权限，普通用户可操作本人来源。"""
    other = _save(test_client, auth_headers, return_body)
    assert test_client.get(BASE + "/list").status_code == 401

    async def load_user() -> CoreUserSchema:
        """真实普通用户的角色数据范围为仅本人。"""
        async with async_db_session() as db:
            user = (await db.execute(select(UserModel).where(UserModel.username == "user"))).scalar_one()
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    auth = AuthSchema(
        user=test_client.portal.call(load_user),
        permissions=[
            *[f"module_inventory:purchase_return:{action}" for action in ["query", "detail", "create", "update", "delete", "confirm", "void", "export"]],
            *[f"module_inventory:purchase_in:{action}" for action in ["create", "confirm"]],
        ],
    )
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        assert test_client.get(BASE + f"/detail/{other['id']}").status_code == 404
        assert test_client.get(BASE + "/list", params={"order_no": other["order_no"]}).json()["data"]["total"] == 0
        assert test_client.post(BASE + "/create", json=return_body).status_code == 404
        assert test_client.post(BASE + f"/update/{other['id']}", json={**return_body, "version": 0}).status_code == 404
        assert test_client.post(BASE + f"/confirm/{other['id']}", json={"version": 0}).status_code == 404
        assert test_client.post(BASE + f"/void/{other['id']}", json={"version": 0, "void_reason": "无权"}).status_code == 404
        assert test_client.post(BASE + "/delete", json={"ids": [other["id"]], "versions": {other["id"]: 0}}).status_code == 404
        exported = test_client.post(BASE + "/export", json={"query": {"order_no": other["order_no"]}, "fields": ["order_no"], "format": "csv"})
        assert list(csv.reader(StringIO(exported.content.decode("utf-8-sig")))) == [["单据编号"]]
        source = _confirm(test_client, {}, _create(test_client, {}, purchase_body))
        own_body = {
            **return_body,
            "source_purchase_in_id": source["id"],
            "items": [
                {
                    "goods_id": source["items"][0]["goods_id"],
                    "source_purchase_in_item_id": source["items"][0]["id"],
                    "quantity": "0.500",
                    "unit_price": source["items"][0]["unit_price"],
                }
            ],
        }
        own = _save(test_client, {}, own_body)
        assert own["created_id"] == auth.user.id
        _post_confirm(test_client, {}, own)
        auth.permissions = []
        assert test_client.get(BASE + "/list").status_code == 403
        assert test_client.post(BASE + f"/confirm/{own['id']}", json={"version": 0}).status_code == 403
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous
