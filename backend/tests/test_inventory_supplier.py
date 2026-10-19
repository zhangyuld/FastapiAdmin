"""供应商空值兼容和电话长度边界，所有数据库写入仅发生在测试库。"""

import csv
from io import BytesIO, StringIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.database import async_db_session
from app.modules.inventory.basic.model import SupplierModel
from app.modules.inventory.basic.supplier.schema import SupplierCreateSchema, SupplierOutSchema, SupplierQueryParam, SupplierUpdateSchema


@pytest.mark.parametrize("schema", [SupplierCreateSchema, SupplierUpdateSchema])
@pytest.mark.parametrize("optional", [{}, {"contact": None, "phone": None}, {"contact": " \t ", "phone": " \n ", "address": "   ", "remark": "  "}])
def test_supplier_optional_fields(schema: type[SupplierCreateSchema], optional: dict[str, Any]) -> None:
    """新增和修改允许省略或传空值，空白可选文本归一为空值。"""
    value = schema(supplier_code="TEST-SUPPLIER-SCHEMA", supplier_name="供应商", status=0, **optional)
    assert value.contact is None and value.phone is None
    assert value.address is None and value.remark is None


@pytest.mark.parametrize("phone", [None, "1" * 21, "1" * 32])
def test_supplier_output_matches_nullable_model(phone: str | None) -> None:
    """直接模拟ORM查询结果，验证空联系人和数据库合法电话均可输出。"""
    supplier = SupplierModel(supplier_code="TEST-SUPPLIER-ORM", supplier_name="供应商", contact=None, phone=phone, status=0, is_deleted=False)
    result = SupplierOutSchema.model_validate(supplier)
    assert result.contact is None and result.phone == phone


@pytest.mark.parametrize("schema", [SupplierCreateSchema, SupplierUpdateSchema, SupplierQueryParam])
def test_supplier_phone_boundary(schema: type[SupplierCreateSchema] | type[SupplierQueryParam]) -> None:
    """请求和查询的电话长度均为32，超过模型上限仍拒绝。"""
    values = {} if schema is SupplierQueryParam else {"supplier_code": "TEST-SUPPLIER-LENGTH", "supplier_name": "供应商", "status": 0}
    assert schema(phone="1" * 32, **values).phone == "1" * 32
    with pytest.raises(ValidationError):
        schema(phone="1" * 33, **values)


def test_supplier_nonempty_optional_text() -> None:
    """正常联系信息保留内容，仅去除首尾空白。"""
    value = SupplierCreateSchema(supplier_code="TEST-SUPPLIER-TEXT", supplier_name="供应商", status=0, contact=" 张三 ", phone=" 13800138000 ", address=" 联系地址 ", remark=" 备注 ")
    assert (value.contact, value.phone, value.address, value.remark) == ("张三", "13800138000", "联系地址", "备注")


@pytest.mark.parametrize("field", ["supplier_code", "supplier_name"])
def test_supplier_required_fields_remain_required(field: str) -> None:
    """联系人可选不影响供应商名称和编码的必填校验。"""
    values = {"supplier_code": "TEST-SUPPLIER-REQUIRED", "supplier_name": "供应商", "status": 0, field: "  "}
    with pytest.raises(ValidationError):
        SupplierCreateSchema(**values)


@pytest.mark.parametrize("phone", [None, "1" * 32])
def test_supplier_database_null_list_detail_and_export(test_client: TestClient, auth_headers: dict[str, str], phone: str | None) -> None:
    """从数据库真实NULL行验证分页、详情以及Excel/CSV导出。"""
    code = f"TEST-SUPPLIER-NULL-{len(phone) if phone else 0}"

    async def insert_supplier() -> int:
        """直接插入允许NULL的ORM数据，覆盖导入脚本和历史数据场景。"""
        async with async_db_session() as db, db.begin():
            supplier = SupplierModel(supplier_code=code, supplier_name="空联系人供应商", contact=None, phone=phone, status=0)
            db.add(supplier)
            await db.flush()
            return supplier.id

    identifier = test_client.portal.call(insert_supplier)
    query = {"supplier_code": code, "page_no": 1, "page_size": 10}
    if phone:
        query["phone"] = phone
    listed = test_client.get("/inventory/basic/supplier/list", headers=auth_headers, params=query)
    assert listed.status_code == 200, listed.text
    items = listed.json()["data"]["items"]
    assert len(items) == 1 and items[0]["id"] == identifier
    assert items[0]["contact"] is None and items[0]["phone"] == phone
    detail = test_client.get(f"/inventory/basic/supplier/detail/{identifier}", headers=auth_headers)
    assert detail.status_code == 200, detail.text
    assert detail.json()["data"]["contact"] is None and detail.json()["data"]["phone"] == phone

    for format_name in ("xlsx", "csv"):
        exported = test_client.post("/inventory/basic/supplier/export", headers=auth_headers, json={"query": {"supplier_code": code}, "fields": ["supplier_code", "contact", "phone"], "format": format_name})
        assert exported.status_code == 200, exported.text
        if format_name == "xlsx":
            workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
                assert rows[0] == ("供应商编码", "联系人", "联系电话")
                assert rows[1] == (code, None, phone)
            finally:
                workbook.close()
        else:
            rows = list(csv.reader(StringIO(exported.content.decode("utf-8-sig"))))
            assert rows == [["供应商编码", "联系人", "联系电话"], [code, "", phone or ""]]


def test_supplier_create_update_optional_text(test_client: TestClient, auth_headers: dict[str, str]) -> None:
    """创建省略联系字段成功，更新清理空白并保留32位电话。"""
    values = {"supplier_code": "TEST-SUPPLIER-OPTIONAL-CRUD", "supplier_name": "联系信息供应商", "status": 0}
    response = test_client.post("/inventory/basic/supplier/create", headers=auth_headers, json=values)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["contact"] is None and data["phone"] is None
    response = test_client.post(f"/inventory/basic/supplier/update/{data['id']}", headers=auth_headers, json={**values, "contact": "   ", "phone": "1" * 32, "address": "   ", "remark": "  "})
    assert response.status_code == 200, response.text
    updated = response.json()["data"]
    assert updated["contact"] is None and updated["phone"] == "1" * 32
    assert updated["address"] is None and updated["remark"] is None
