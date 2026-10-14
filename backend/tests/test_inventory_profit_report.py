"""毛利报表优惠收入、固化成本、期间退货、比例与独立授权回归。"""

from datetime import date
from decimal import Decimal
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import pytest
from sqlalchemy import select

from app.core.database import async_db_session
from app.modules.inventory.basic.model import CustomerModel, GoodsCategoryModel, GoodsModel, WarehouseModel
from app.modules.inventory.sale.model import SaleOutItemModel, SaleOutModel, SaleReturnItemModel, SaleReturnModel
from app.modules.system.user.model import UserModel

PREFIX = "/inventory/report/profit"


@pytest.fixture
def profit_data(test_client, auth_headers):
    """只在临时库构造已固化成本，包含优惠、多行、跨月退货与无效单。"""
    key = uuid4().hex[:12]

    async def seed():
        async with async_db_session() as db, db.begin():
            user_id = await db.scalar(select(UserModel.id).where(UserModel.username == "super"))
            customer = CustomerModel(customer_code=key, customer_name=key + "Customer", status=1, created_id=user_id)
            categories = [GoodsCategoryModel(name=key + "Category" + str(i)) for i in range(2)]
            warehouses = [WarehouseModel(warehouse_code=key + str(i), warehouse_name=key + "Warehouse" + str(i), status=1) for i in range(2)]
            db.add_all([customer, *categories, *warehouses])
            await db.flush()
            goods = [
                GoodsModel(
                    goods_code=key + ("G0" if i == 0 else "%_G1"), goods_name=key + "Goods" + str(i), category_id=categories[i].id, unit="件" if i == 0 else "箱", reference_cost_price=Decimal("999")
                )
                for i in range(2)
            ]
            db.add_all(goods)
            await db.flush()

            async def document(is_return, day, warehouse, lines, discount="0", status=1, source=None, deleted=False):
                """销售明细保存原价与独立成本，退货保存已分摊退款和原成本。"""
                total = sum((Decimal(amount) for _, _, amount, _, _ in lines), Decimal("0"))
                values = {
                    "order_no": key + ("SR" if is_return else "SO") + uuid4().hex[:10],
                    "customer_id": customer.id,
                    "warehouse_id": warehouses[warehouse].id,
                    "business_date": date.fromisoformat(day),
                    "status": status,
                    "created_id": user_id,
                    "is_deleted": deleted,
                }
                if is_return:
                    values.update(return_amount=total, source_sale_out_id=source[0].id, return_reason="report-test")
                else:
                    values.update(original_amount=total, discount_amount=Decimal(discount), receivable_amount=total - Decimal(discount))
                row = (SaleReturnModel if is_return else SaleOutModel)(**values)
                db.add(row)
                await db.flush()
                created = []
                for index, (g, quantity, amount, cost, unit) in enumerate(lines, 1):
                    q, c = Decimal(quantity), Decimal(cost) if cost is not None else None
                    values = {
                        "line_no": index,
                        "goods_id": goods[g].id,
                        "goods_code": goods[g].goods_code,
                        "goods_name": goods[g].goods_name,
                        "spec": "snapshot",
                        "unit": unit,
                        "quantity": q,
                        "unit_price": Decimal(amount) / q,
                        "line_amount": Decimal(amount),
                        "unit_cost": c / q if c is not None else None,
                        "cost_amount": c,
                        "created_id": user_id,
                    }
                    if is_return:
                        values.update(sale_return_id=row.id, source_sale_out_item_id=source[1][0].id)
                    else:
                        values["sale_out_id"] = row.id
                    item = (SaleReturnItemModel if is_return else SaleOutItemModel)(**values)
                    db.add(item)
                    created.append(item)
                await db.flush()
                return row, created

            original = await document(False, "2026-09-30", 0, [(0, "10", "100", "40", "件"), (1, "2", "200", "100", "箱"), (0, "2", "20", "8", "件")], discount="32")
            await document(False, "2026-10-05", 0, [(0, "3", "30", "9", "箱")])
            warehouse_original = await document(False, "2026-10-06", 1, [(0, "1", "10", "5", "件")])
            returned = await document(True, "2026-10-10", 0, [(0, "3", "27", "12", "件")], source=original)
            await document(True, "2026-10-11", 1, [(0, "1", "10", "5", "件")], source=warehouse_original)
            for status in [0, 2]:
                await document(False, "2026-10-05", 0, [(0, "1", "999", None, "件")], status=status)
                await document(True, "2026-10-12", 0, [(0, "1", "10", "4", "件")], status=status, source=original)
            await document(False, "2026-10-05", 0, [(0, "1", "777", None, "件")], deleted=True)
            _, deleted_items = await document(False, "2026-10-05", 0, [(0, "1", "888", None, "件")])
            deleted_items[0].is_deleted = True
            goods[0].goods_code = key + "CURRENT"
            goods[0].goods_name = "renamed"
            customer.is_deleted = True
            return {
                "key": key,
                "customer_name": customer.customer_name,
                "customer_id": customer.id,
                "warehouse_ids": [row.id for row in warehouses],
                "category_ids": [row.id for row in categories],
                "goods_ids": [row.id for row in goods],
                "goods_codes": [key + "G0", key + "%_G1"],
                "original_id": original[0].id,
                "return_id": returned[0].id,
                "item_ids": [row.id for row in original[1]],
                "business_no": original[0].order_no,
            }

    with patch("app.core.router_class._write_operation_log_async", new=AsyncMock()), patch("app.modules.system.auth.service._write_login_log", new=AsyncMock()):
        yield test_client.portal.call(seed)


def data(client, headers, fixture, endpoint, params=None):
    """报表断言限定独立测试客户，不串入其他用例的经营数据。"""
    response = client.get(PREFIX + endpoint, headers=headers, params={"customer_name": fixture["customer_name"], **(params or {})})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_profit_summary_and_source_identity(test_client, auth_headers, profit_data):
    """优惠后收入扣退款，成本扣原成本退回，汇总不平均明细比例。"""
    summary = data(test_client, auth_headers, profit_data, "/summary")
    assert summary == {"net_sale_amount": "291.00", "net_cost_amount": "145.00", "gross_profit": "146.00", "gross_margin": "50.17"}
    rows = data(test_client, auth_headers, profit_data, "/list")
    assert rows["total"] == 5 and len({row["id"] for row in rows["items"]}) == 5
    returned = next(row for row in rows["items"] if row["source"]["business_id"] == profit_data["return_id"] and row["source"]["business_type"] == "sale_return")
    assert (returned["net_sale_amount"], returned["net_cost_amount"], returned["gross_profit"], returned["gross_margin"]) == ("-27.00", "-12.00", "-15.00", None)
    assert all("margin_order" not in row and "cost_missing" not in row for row in rows["items"])
    sale = test_client.get("/inventory/report/sale/summary", headers=auth_headers, params={"customer_name": profit_data["customer_name"]}).json()["data"]
    assert summary["net_sale_amount"] == sale["net_sale_amount"]


def test_cross_month_return_and_margin_over_100(test_client, auth_headers, profit_data):
    """退货使用自身月份，成本退回大于当期出库成本时允许超过百分之百。"""
    september = data(test_client, auth_headers, profit_data, "/summary", {"business_date": ["2026-09-01", "2026-09-30"]})
    assert september == {"net_sale_amount": "288.00", "net_cost_amount": "148.00", "gross_profit": "140.00", "gross_margin": "48.61"}
    october = data(test_client, auth_headers, profit_data, "/summary", {"business_date": ["2026-10-01", "2026-10-31"]})
    assert october == {"net_sale_amount": "3.00", "net_cost_amount": "-3.00", "gross_profit": "6.00", "gross_margin": "200.00"}
    returned = data(test_client, auth_headers, profit_data, "/summary", {"business_date": ["2026-10-10", "2026-10-10"]})
    assert returned == {"net_sale_amount": "-27.00", "net_cost_amount": "-12.00", "gross_profit": "-15.00", "gross_margin": None}


def test_goods_snapshot_filter_and_full_discount_denominator(test_client, auth_headers, profit_data):
    """商品筛选发生在整单分摊后，不带入其他商品收入或成本，快照不跟随改名。"""
    result = data(test_client, auth_headers, profit_data, "/summary", {"goods_code": profit_data["goods_codes"][0]})
    assert result == {"net_sale_amount": "111.00", "net_cost_amount": "45.00", "gross_profit": "66.00", "gross_margin": "59.46"}
    query = {"business_no": profit_data["business_no"], "goods_code": profit_data["goods_codes"][0]}
    row = data(test_client, auth_headers, profit_data, "/list", query)["items"][0]
    assert row["net_sale_amount"] == "108.00" and row["net_cost_amount"] == "48.00"
    literal = data(test_client, auth_headers, profit_data, "/summary", {"goods_code": "%_"})
    assert literal["net_sale_amount"] == "180.00" and literal["net_cost_amount"] == "100.00"
    category = data(test_client, auth_headers, profit_data, "/summary", {"category_id": profit_data["category_ids"][1]})
    assert category == literal
    assert data(test_client, auth_headers, profit_data, "/summary", {"category_name": profit_data["key"] + "Category1"}) == literal


def test_groups_units_stable_paging_and_margin_sort(test_client, auth_headers, profit_data):
    """不同单位和仓库不合并，净零组仍保留，稳定 ID 不依赖分页顺序。"""
    result = data(test_client, auth_headers, profit_data, "/list", {"view": "goods"})
    assert result["total"] == 4 and all("source" not in row for row in result["items"])
    row = next(row for row in result["items"] if row["goods_code"] == profit_data["goods_codes"][0] and row["unit"] == "件" and row["warehouse_name"].endswith("0"))
    assert (row["net_sale_amount"], row["net_cost_amount"], row["gross_profit"], row["gross_margin"]) == ("81.00", "36.00", "45.00", "55.56")
    zero = next(row for row in result["items"] if row["warehouse_name"].endswith("1"))
    assert zero["gross_profit"] == "0.00" and zero["gross_margin"] is None
    ids = [data(test_client, auth_headers, profit_data, "/list", {"view": "goods", "page_no": i, "page_size": 1})["items"][0]["id"] for i in range(1, 5)]
    assert set(ids) == {row["id"] for row in result["items"]} and len(ids) == len(set(ids))
    assert data(test_client, auth_headers, profit_data, "/summary", {"view": "goods", "page_size": 1})["gross_margin"] == "50.17"
    ordered = data(test_client, auth_headers, profit_data, "/list", {"order_by": '[{"gross_margin":"desc"}]'})["items"]
    assert [row["gross_margin"] for row in ordered] == ["70.00", "50.00", "48.61", None, None]


def test_reference_and_current_inventory_cost_do_not_change_profit(test_client, auth_headers, profit_data):
    """档案参考进价和当前库存成本改变，不影响原销售固化成本。"""
    from app.modules.inventory.stock.model import StockBalanceModel

    async def change():
        async with async_db_session() as db, db.begin():
            goods = await db.get(GoodsModel, profit_data["goods_ids"][0])
            goods.reference_cost_price = Decimal("10000")
            db.add(StockBalanceModel(goods_id=goods.id, warehouse_id=profit_data["warehouse_ids"][0], quantity=999, avg_cost=900, stock_amount=899100))

    before = data(test_client, auth_headers, profit_data, "/summary")
    test_client.portal.call(change)
    assert data(test_client, auth_headers, profit_data, "/summary") == before


@pytest.mark.parametrize("cost,expected_profit,expected_margin", [("0", "288.00", "100.00"), ("400", "-912.00", "-316.67")])
def test_zero_cost_and_loss_are_not_missing_cost(test_client, auth_headers, profit_data, cost, expected_profit, expected_margin):
    """真实零成本可计算毛利，低于成本销售允许负毛利和负毛利率。"""

    async def change():
        async with async_db_session() as db, db.begin():
            for id in profit_data["item_ids"]:
                item = await db.get(SaleOutItemModel, id)
                item.cost_amount = Decimal(cost)

    test_client.portal.call(change)
    query = {"business_date": ["2026-09-30", "2026-09-30"]}
    summary = data(test_client, auth_headers, profit_data, "/summary", query)
    assert summary["gross_profit"] == expected_profit and summary["gross_margin"] == expected_margin


def test_full_discount_and_zero_income_preserve_cost(test_client, auth_headers, profit_data):
    """全额优惠仍保留固化成本与负毛利，零净收入比例为空而不是零。"""

    async def change():
        async with async_db_session() as db, db.begin():
            original = await db.get(SaleOutModel, profit_data["original_id"])
            original.discount_amount = original.original_amount
            original.receivable_amount = Decimal("0")

    test_client.portal.call(change)
    summary = data(test_client, auth_headers, profit_data, "/summary", {"business_date": ["2026-09-30", "2026-09-30"]})
    assert summary == {"net_sale_amount": "0.00", "net_cost_amount": "148.00", "gross_profit": "-148.00", "gross_margin": None}


def test_missing_cost_rejects_full_scope_but_not_unrelated_goods(test_client, auth_headers, profit_data):
    """不能只检查当前页，命中缺失成本时列表、汇总、导出都拒绝且不静默置零。"""

    async def change():
        async with async_db_session() as db, db.begin():
            row = await db.get(SaleOutItemModel, profit_data["item_ids"][1])
            row.cost_amount = None

    test_client.portal.call(change)
    for endpoint in ["/list", "/summary"]:
        response = test_client.get(PREFIX + endpoint, headers=auth_headers, params={"customer_name": profit_data["customer_name"], "page_size": 1})
        assert response.status_code == 409 and "固化成本" in response.text
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json={"query": {"customer_name": profit_data["customer_name"]}, "fields": ["gross_profit"]})
    assert response.status_code == 409
    unaffected = data(test_client, auth_headers, profit_data, "/summary", {"goods_code": profit_data["goods_codes"][0]})
    assert unaffected["gross_profit"] == "66.00"
    assert data(test_client, auth_headers, profit_data, "/summary", {"business_date": ["2026-10-01", "2026-10-31"]})["gross_margin"] == "200.00"


def test_void_recalculates_income_and_cost(test_client, auth_headers, profit_data):
    """作废移除原业务的收入和成本，不再读取冲销流水重复调整毛利。"""

    async def void_return():
        async with async_db_session() as db, db.begin():
            returned = await db.get(SaleReturnModel, profit_data["return_id"])
            returned.status = 2

    test_client.portal.call(void_return)
    summary = data(test_client, auth_headers, profit_data, "/summary")
    assert summary["net_sale_amount"] == "318.00" and summary["net_cost_amount"] == "157.00" and summary["gross_profit"] == "161.00"

    async def void_original():
        async with async_db_session() as db, db.begin():
            original = await db.get(SaleOutModel, profit_data["original_id"])
            original.status = 2

    test_client.portal.call(void_original)
    summary = data(test_client, auth_headers, profit_data, "/summary")
    assert summary["net_sale_amount"] == "30.00" and summary["net_cost_amount"] == "9.00" and summary["gross_margin"] == "70.00"


@pytest.mark.parametrize(
    "params",
    [
        {"view": "movement"},
        {"warehouse_id": 0},
        {"business_date": ["2026-10-02", "2026-10-01"]},
        {"business_date": ["2026-10-01"]},
        {"view": "goods", "business_no": "SO"},
        {"order_by": '[{"margin_order":"asc"}]'},
        {"order_by": '[{"gross_profit":"drop"}]'},
        {"order_by": "[{}]"},
        {"view": "goods", "order_by": '[{"snapshot_unit":"asc"}]'},
    ],
)
def test_invalid_query(test_client, auth_headers, profit_data, params):
    """视图、日期、跨视图单号与内部排序字段由后端校验。"""
    assert test_client.get(PREFIX + "/list", headers=auth_headers, params={"customer_name": profit_data["customer_name"], **params}).status_code == 422


def test_empty_results_keep_null_margin(test_client, auth_headers, profit_data):
    """空数据金额为真实零，毛利率不可计算返回 null，不伪造百分之零。"""
    query = {"goods_code": "unmatched"}
    assert data(test_client, auth_headers, profit_data, "/summary", query) == {"net_sale_amount": "0.00", "net_cost_amount": "0.00", "gross_profit": "0.00", "gross_margin": None}
    assert data(test_client, auth_headers, profit_data, "/list", query)["items"] == []


def test_exports_follow_field_order_percent_and_null_placeholder(test_client, auth_headers, profit_data):
    """导出百分号与不可计算占位和页面一致，字段选择不导出其他隐藏指标。"""
    from io import BytesIO

    from openpyxl import load_workbook

    body = {
        "query": {"customer_name": profit_data["customer_name"]},
        "fields": ["gross_margin", "gross_profit", "business_no", "gross_margin"],
        "format": "csv",
        "filename": "毛利核对.xlsx",
    }
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "毛利率,毛利,来源单号" and len(lines) == 6
    assert any(line.startswith("48.61%,140.00,") for line in lines[1:])
    assert any(line.startswith("—,-15.00,") for line in lines[1:])
    assert "%E6%AF%9B%E5%88%A9%E6%A0%B8%E5%AF%B9.csv" in response.headers["content-disposition"]
    body.update(query={"view": "goods", "customer_name": profit_data["customer_name"]}, fields=["goods_name", "gross_margin", "net_cost_amount"], format="xlsx", filename="", sheet_name="毛利核对")
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "毛利核对" and sheet.max_row == 5 and sheet.max_column == 3
    assert [cell.value for cell in sheet[1]] == ["商品名称", "毛利率", "净销售成本"]
    assert any(row[1].value == "—" for row in list(sheet.rows)[1:])
    assert "%E6%AF%9B%E5%88%A9%E6%8A%A5%E8%A1%A8_" in response.headers["content-disposition"]
    body["fields"] = ["business_no"]
    assert test_client.post(PREFIX + "/export", headers=auth_headers, json=body).status_code == 422


def test_empty_export_has_selected_header(test_client, auth_headers, profit_data):
    """空结果生成合法表头文件，不将无数据误判为未知成本或零毛利率。"""
    response = test_client.post(
        PREFIX + "/export",
        headers=auth_headers,
        json={
            "query": {"customer_name": profit_data["customer_name"], "goods_code": "unmatched"},
            "fields": ["gross_margin"],
            "format": "csv",
        },
    )
    assert response.status_code == 200, response.text
    assert response.content.decode("utf-8-sig").strip() == "毛利率"


@pytest.mark.parametrize(
    "options",
    [
        {"fields": []},
        {"fields": ["sale_amount"]},
        {"format": "pdf"},
        {"filename": "../bad"},
        {"filename": "a\nb"},
        {"sheet_name": "bad[]"},
        {"query": {"view": "balance"}},
        {"query": {"view": "goods", "business_no": "SO"}},
        {"unexpected": True},
    ],
)
def test_invalid_export_options(test_client, auth_headers, options):
    """服务端独立验证格式、文件名、字段和视图，而不是信任前端选项。"""
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json={"fields": ["gross_profit"], **options})
    assert response.status_code == 422, response.text


def test_formula_text_guard_keeps_negative_profit_and_percentage(test_client, auth_headers, profit_data):
    """档案文本按字面值导出，负毛利与服务端百分数不添加文本防护前缀。"""
    import csv
    from io import BytesIO, StringIO

    from openpyxl import load_workbook

    async def rename():
        async with async_db_session() as db, db.begin():
            row = await db.get(CustomerModel, profit_data["customer_id"])
            row.customer_name = "=2+2"

    test_client.portal.call(rename)
    body = {"query": {"warehouse_name": profit_data["key"]}, "fields": ["customer_name", "gross_profit", "gross_margin"], "format": "csv"}
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
    assert all(row[0] == "'=2+2" for row in rows[1:])
    assert any(row[1] == "-15.00" and row[2] == "—" for row in rows[1:])
    assert any(row[2] == "48.61%" for row in rows[1:])
    body["format"] = "xlsx"
    response = test_client.post(PREFIX + "/export", headers=auth_headers, json=body)
    cell = load_workbook(BytesIO(response.content)).active.cell(2, 1)
    assert cell.value == "'=2+2" and cell.data_type == "s"


def test_export_over_limit_is_explicit(test_client):
    """超过十万行明确拒绝，不能让文件工具静默截断完整毛利结果。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.exceptions import CustomException
    from app.modules.inventory.report.profit.schema import ProfitReportExport
    from app.modules.inventory.report.profit.service import ProfitReportService

    async def run():
        async with async_db_session() as db:
            auth = AuthSchema(user=CoreUserSchema(id=1, username="super", name="super", dept_id=None, is_superuser=True))
            service = ProfitReportService(auth, db)
            with patch.object(service.crud, "rows", new=AsyncMock(return_value=([], 100001))):
                with pytest.raises(CustomException, match="10万行"):
                    await service.export(ProfitReportExport(fields=["gross_profit"]))

    test_client.portal.call(run)


def test_independent_permissions_readonly_and_source_scope(test_client, auth_headers, profit_data):
    """毛利查询覆盖企业业务，但不继承原销售详情数据范围、导出或其他报表权限。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.dependencies import get_current_user

    async def load_user():
        async with async_db_session() as db:
            user = await db.scalar(select(UserModel).where(UserModel.username == "user"))
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    auth = AuthSchema(user=test_client.portal.call(load_user), permissions=["module_inventory:profit_report:query", "module_inventory:sale_out:detail"])
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        assert data(test_client, {}, profit_data, "/summary")["gross_profit"] == "146.00"
        assert test_client.get("/inventory/sale/out/detail/" + str(profit_data["original_id"])).status_code == 404
        assert test_client.post(PREFIX + "/export", json={"fields": ["gross_profit"]}).status_code == 403
        assert test_client.get("/inventory/report/sale/list").status_code == 403
        auth.permissions = ["module_inventory:sale_report:query", "module_inventory:sale_out:query"]
        assert test_client.get(PREFIX + "/list").status_code == 403
        assert test_client.get(PREFIX + "/summary").status_code == 403
        auth.permissions = ["module_inventory:profit_report:export"]
        assert test_client.post(PREFIX + "/export", json={"query": {"customer_name": profit_data["customer_name"]}, "fields": ["gross_profit"]}).status_code == 200
        assert test_client.get(PREFIX + "/list").status_code == 403
        for endpoint in ["create", "update/1", "delete", "confirm/1", "void/1"]:
            assert test_client.post(PREFIX + "/" + endpoint, json={}).status_code in (404, 405)
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous


def test_mysql_grouping_keeps_binary_snapshot_keys():
    """MySQL 方言按完整快照二进制键分组，内部排序比例不暴露为业务字段。"""
    from sqlalchemy.dialects import mysql

    from app.modules.inventory.report.profit.crud import ProfitReportCRUD
    from app.modules.inventory.report.profit.schema import ProfitReportQuery

    lines = ProfitReportCRUD._branch(ProfitReportQuery(), False).subquery()
    grouped = ProfitReportCRUD._grouped(lines, "goods")
    sql = str(select(grouped).compile(dialect=mysql.dialect()))
    assert "AS BINARY" in sql and "snapshot_unit" in sql and "CAST" in sql.split("GROUP BY", 1)[1]


@pytest.mark.parametrize(
    "sale,cost,profit,margin",
    [
        ("3.00", "1.00", "2.00", "66.67"),
        ("8.00", "7.01", "0.99", "12.38"),
        ("8.00", "8.99", "-0.99", "-12.38"),
        ("3.00", "-3.00", "6.00", "200.00"),
        ("0.00", "5.00", "-5.00", None),
        ("-27.00", "-12.00", "-15.00", None),
        ("1000.00", "1000.01", "-0.01", "0.00"),
        ("18000000000000000.00", "9000000000000000.00", "9000000000000000.00", "50.00"),
    ],
)
def test_decimal_margin_rounding_and_aggregate_precision(sale, cost, profit, margin):
    """Decimal HALF_UP 处理正负半分与聚合大金额，消除负零百分数。"""
    from app.modules.inventory.report.profit.schema import ProfitReportSummary
    from app.modules.inventory.report.profit.service import profit_metrics

    result = ProfitReportSummary.model_validate(profit_metrics(Decimal(sale), Decimal(cost))).model_dump(mode="json")
    assert result == {"net_sale_amount": sale, "net_cost_amount": cost, "gross_profit": profit, "gross_margin": margin}
