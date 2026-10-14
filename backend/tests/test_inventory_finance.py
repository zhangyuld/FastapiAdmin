"""往来查询接口回归：方向金额、冲销关联、分页和导出。"""

from datetime import date
from decimal import Decimal
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import pytest
from sqlalchemy import select

from app.core.database import async_db_session
from app.modules.inventory.basic.model import CustomerModel, SupplierModel
from app.modules.inventory.finance.model import PayableModel, ReceivableModel
from app.modules.system.user.model import UserModel


@pytest.fixture(params=["receivable", "payable"])
def ledger(request, test_client, auth_headers):
    """在测试数据库建立两方余额及一组冲销，不写入开发数据库。"""
    kind = request.param
    code = uuid4().hex[:12]

    async def seed():
        async with async_db_session() as db, db.begin():
            user = await db.scalar(select(UserModel.id).limit(1))
            party_model = CustomerModel if kind == "receivable" else SupplierModel
            prefix = "customer" if kind == "receivable" else "supplier"
            parties = [party_model(**{prefix + "_code": code + str(n), prefix + "_name": "Finance" + str(n)}, created_id=user) for n in range(2)]
            db.add_all(parties)
            await db.flush()
            model = ReceivableModel if kind == "receivable" else PayableModel
            rows = []
            for n, (p, amount, allocated, operation) in enumerate(
                [(0, "100", "20", "confirm"), (0, "-30", "10", "confirm"), (1, "-40", "0", "confirm"), (0, "50", "0", "confirm"), (0, "-50", "0", "void")]
            ):
                row = model(
                    **{prefix + "_id": parties[p].id},
                    business_type="sale_out" if kind == "receivable" else "purchase_in",
                    business_id=parties[0].id * 100 + (3 if operation == "void" else n),
                    business_no=code + str(n),
                    operation_type=operation,
                    amount=Decimal(amount),
                    allocated_amount=Decimal(allocated),
                    unallocated_amount=abs(Decimal(amount)) - Decimal(allocated),
                    bill_date=date(2026, 10, 13),
                    status=1 if allocated != "0" else 0,
                    created_id=user,
                )
                if operation == "void":
                    row.reversal_id = rows[3].id
                db.add(row)
                await db.flush()
                rows.append(row)
            return {"kind": kind, "code": code, "parties": [p.id for p in parties], "rows": [r.id for r in rows]}

    with patch("app.core.router_class._write_operation_log_async", new=AsyncMock()), patch("app.modules.system.auth.service._write_login_log", new=AsyncMock()):
        yield test_client.portal.call(seed)


def get_data(client, headers, ledger, path, params=None):
    """读取统一响应并保留失败响应便于定位。"""
    response = client.get("/inventory/finance/" + ledger["kind"] + path, headers=headers, params=params)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_directional_balance(test_client, auth_headers, ledger):
    """原单与冲销抵消，退货按负向未核销金额计入。"""
    data = get_data(test_client, auth_headers, ledger, "/balance/list", {"party_id": ledger["parties"][0]})
    row = data["items"][0]
    assert row["business_net_amount"] == "70.00"
    assert row["unsettled_net_amount"] == "60.00"
    assert row["allocated_net_amount"] == "10.00"
    assert row["bill_count"] == 4


def test_summary_groups_before_sign(test_client, auth_headers, ledger):
    """正负双方独立汇总，不以跨对象净额代替待收待付。"""
    data = get_data(test_client, auth_headers, ledger, "/summary", {"party_code": ledger["code"]})
    assert data["party_count"] == 2
    assert data["positive_balance"] == "60.00"
    assert data["negative_balance"] == "40.00"
    assert data["unsettled_net_amount"] == "20.00"


def test_reversal_detail(test_client, auth_headers, ledger):
    """冲销不伪装核销，原单与反向记录均可关联追溯。"""
    for position, state, related in [(3, "reversed", 4), (4, "reversal", 3)]:
        data = get_data(test_client, auth_headers, ledger, "/detail/" + str(ledger["rows"][position]))
        assert data["reversal_state"] == state
        assert data["status"] == 0
        assert [r["id"] for r in data["related_bills"]] == [ledger["rows"][related]]


def test_filters_and_paging(test_client, auth_headers, ledger):
    """账单筛选和余额符号筛选不受当前页影响。"""
    rows = get_data(test_client, auth_headers, ledger, "/list", {"party_code": ledger["code"], "reversal_state": "reversed", "page_size": 1})
    assert rows["total"] == 1
    assert rows["items"][0]["id"] == ledger["rows"][3]
    rows = get_data(test_client, auth_headers, ledger, "/balance/list", {"party_code": ledger["code"], "balance_state": "negative"})
    assert rows["total"] == 1
    assert rows["items"][0]["party_id"] == ledger["parties"][1]


def test_csv_export_and_validation(test_client, auth_headers, ledger):
    """全量导出使用选择的字段顺序并校验非法字段。"""
    path = "/inventory/finance/" + ledger["kind"] + "/export"
    body = {"query": {"party_code": ledger["code"]}, "fields": ["business_no", "amount"], "format": "csv", "filename": "finance"}
    response = test_client.post(path, headers=auth_headers, json=body)
    assert response.status_code == 200, response.text
    assert "finance.csv" in response.headers["content-disposition"]
    lines = response.content.decode("utf-8-sig").splitlines()
    assert len(lines) == 6
    assert len(lines[0].split(",")) == 2
    body["fields"] = ["unknown"]
    assert test_client.post(path, headers=auth_headers, json=body).status_code == 422


def test_readonly_and_missing(test_client, auth_headers, ledger):
    """往来查询不暴露手工修改账单入口。"""
    prefix = "/inventory/finance/" + ledger["kind"]
    assert test_client.get(prefix + "/detail/2147483647", headers=auth_headers).status_code == 404
    for endpoint in ["create", "update/1", "delete"]:
        assert test_client.post(prefix + "/" + endpoint, headers=auth_headers, json={}).status_code in (404, 405)


def test_shared_finance_and_button_permissions(test_client, auth_headers, ledger):
    """财务权限允许查询完整往来，不由创建人截断余额，详情与导出独立授权。"""
    from app.core.base_schema import AuthSchema, CoreUserSchema
    from app.core.dependencies import get_current_user

    async def load_user():
        async with async_db_session() as db:
            user = await db.scalar(select(UserModel).where(UserModel.username == "user"))
            return CoreUserSchema(id=user.id, username=user.username, name=user.name, dept_id=user.dept_id, is_superuser=False)

    prefix = "module_inventory:" + ledger["kind"]
    auth = AuthSchema(user=test_client.portal.call(load_user), permissions=[prefix + ":query"])
    previous = test_client.app.dependency_overrides.get(get_current_user)
    test_client.app.dependency_overrides[get_current_user] = lambda: auth
    try:
        data = get_data(test_client, {}, ledger, "/balance/list", {"party_code": ledger["code"]})
        assert data["total"] == 2
        path = "/inventory/finance/" + ledger["kind"]
        assert test_client.get(path + "/detail/" + str(ledger["rows"][0])).status_code == 403
        assert test_client.post(path + "/export", json={"fields": ["amount"]}).status_code == 403
        auth.permissions = []
        for endpoint in ["/list", "/balance/list", "/summary"]:
            assert test_client.get(path + endpoint).status_code == 403
    finally:
        if previous is None:
            test_client.app.dependency_overrides.pop(get_current_user, None)
        else:
            test_client.app.dependency_overrides[get_current_user] = previous


def test_balance_export_and_invalid_range(test_client, auth_headers, ledger):
    """余额全量导出不受分页影响，账单日期校验拒绝逆序。"""
    from io import BytesIO

    from openpyxl import load_workbook

    path = "/inventory/finance/" + ledger["kind"]
    body = {"query": {"party_code": ledger["code"]}, "fields": ["party_name", "unsettled_net_amount"], "format": "xlsx", "filename": ""}
    response = test_client.post(path + "/balance/export", headers=auth_headers, json=body)
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith(".xlsx")
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 3 and sheet.max_column == 2
    response = test_client.get(path + "/list", headers=auth_headers, params=[("bill_date", "2026-10-14"), ("bill_date", "2026-10-13")])
    assert response.status_code == 422
